import calendar
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Optional

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from ..database.db import get_db
from ..models.blinkit_report_store_metadata import BlinkitReportStoreMetadata
from ..models.order import Order, OrderStatusEnum
from ..models.partner_store import PartnerStore
from ..models.store import Store
from ..models.user import RoleEnum, User
from ..schemas.blinkit_report import BlinkitMetadataImportOut, BlinkitMonthlyReportOut
from ..utils.jwt import get_current_user
from ..utils.permissions import get_employee_permissions
from .delivery_reports import aggregate_effective_delivered_by_ist_date, ist_date_range_utc


router = APIRouter(prefix="/blinkit-reports", tags=["Blinkit Reports"])

INCLUDED_STATUSES = (
    OrderStatusEnum.delivered_confirmed,
    OrderStatusEnum.delivered_pending_confirmation,
)
MIN_REPORT_YEAR = 2020
MAX_REPORT_YEAR = 2100
BLINKIT_MODE_OF_OPS = "Real Estate"
BLINKIT_VENDOR_NAME = "VEE KAY AQUATECH Private Limited"
BASE_COLUMNS = [
    "S.No",
    "Region",
    "Entity",
    "State",
    "city",
    "Citi Lead",
    "Outlet Names",
    "Outlet Id",
    "Mode of Ops",
    "Live Date",
    "Supply start date",
    "Status",
    "Closed/Dropped Date",
    "Admin POC",
    "Contact No",
    "Vendor Name",
]
DATE_COLUMNS = {"Live Date", "Supply start date", "Closed/Dropped Date"}
IMPORT_COLUMNS = {
    "Outlet Id",
    "Citi Lead",
    "Live Date",
    "Supply start date",
    "Status",
    "Vendor Rate",
}


def _require_access(current_user: User, permission_key: str, db: Session) -> None:
    if current_user.role == RoleEnum.superadmin:
        return
    if current_user.role != RoleEnum.employee:
        raise HTTPException(status_code=403, detail="Not authorized")
    permission_keys = set(get_employee_permissions(current_user.id, db=db))
    if not ({"*", "*.*"} & permission_keys) and permission_key not in permission_keys:
        raise HTTPException(status_code=403, detail="Permission denied")


def _can_view_vendor_rate(current_user: User, db: Session) -> bool:
    if current_user.role == RoleEnum.superadmin:
        return True
    if current_user.role != RoleEnum.employee:
        return False
    permission_keys = set(get_employee_permissions(current_user.id, db=db))
    return bool(
        {"*", "*.*", "blinkit_reports.vendor_rate_export"} & permission_keys
    )


def month_dates(year: int, month: int) -> list[date]:
    if not MIN_REPORT_YEAR <= year <= MAX_REPORT_YEAR:
        raise HTTPException(
            status_code=400,
            detail=f"Year must be between {MIN_REPORT_YEAR} and {MAX_REPORT_YEAR}",
        )
    if not 1 <= month <= 12:
        raise HTTPException(status_code=400, detail="Month must be between 1 and 12")
    try:
        day_count = calendar.monthrange(year, month)[1]
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="Invalid year or month")
    return [date(year, month, day) for day in range(1, day_count + 1)]


def day_labels(days: list[date]) -> list[str]:
    return [f"{day.day}-{calendar.month_abbr[day.month]}" for day in days]


def _store_query(
    db: Session,
    *,
    region: Optional[str],
    state_filter: Optional[str],
    city: Optional[str],
    entity: Optional[str],
    store_id: Optional[str],
    status_filter: Optional[str],
    vendor_name: Optional[str],
):
    query = (
        db.query(Store, BlinkitReportStoreMetadata)
        .options(
            joinedload(Store.partner_associations).joinedload(PartnerStore.user)
        )
        .outerjoin(
            BlinkitReportStoreMetadata,
            BlinkitReportStoreMetadata.store_id == Store.id,
        )
        .filter(func.upper(Store.channel) == "BLINKIT")
    )
    filters = (
        (region, Store.region),
        (state_filter, Store.state),
        (city, Store.city),
        (entity, Store.entity),
        (status_filter, BlinkitReportStoreMetadata.report_status),
    )
    for value, column in filters:
        if value:
            query = query.filter(func.lower(column) == value.strip().lower())
    if store_id:
        query = query.filter(Store.id == store_id.strip())
    if vendor_name and vendor_name.strip().casefold() != BLINKIT_VENDOR_NAME.casefold():
        query = query.filter(False)
    return query.order_by(
        Store.region.asc().nulls_last(),
        Store.state.asc().nulls_last(),
        Store.city.asc().nulls_last(),
        Store.id.asc(),
    )


def resolve_blinkit_poc(
    store: Store,
    metadata: Optional[BlinkitReportStoreMetadata],
) -> tuple[str, str]:
    mapped_users = [
        association.user
        for association in (store.partner_associations or [])
        if association.user is not None
    ]
    mapped_users.sort(
        key=lambda user: (
            user.role != RoleEnum.partner,
            user.id if user.id is not None else float("inf"),
        )
    )
    mapped_user = mapped_users[0] if mapped_users else None

    raw_metadata_poc = str(metadata.admin_poc or "") if metadata else ""
    raw_metadata_contact = str(metadata.contact_no or "") if metadata else ""
    metadata_poc = raw_metadata_poc if raw_metadata_poc.strip() else ""
    metadata_contact = raw_metadata_contact if raw_metadata_contact.strip() else ""
    admin_poc = metadata_poc or (
        str(mapped_user.full_name or "") if mapped_user else ""
    )
    contact_no = metadata_contact or (
        str(mapped_user.mobile_number or "") if mapped_user else ""
    )
    return admin_poc, contact_no


def _monthly_totals(db: Session, store_ids: list[str], days: list[date]):
    if not store_ids:
        return {}
    start_utc, end_utc = ist_date_range_utc(days[0], days[-1])
    orders = (
        db.query(Order)
        .filter(
            Order.store_id.in_(store_ids),
            Order.order_date >= start_utc,
            Order.order_date < end_utc,
            Order.status.in_(INCLUDED_STATUSES),
        )
        .all()
    )
    orders_by_store = {}
    for order in orders:
        orders_by_store.setdefault(order.store_id, []).append(order)
    return {
        store_id: aggregate_effective_delivered_by_ist_date(store_orders)
        for store_id, store_orders in orders_by_store.items()
    }


def _row(
    serial_number: int,
    store: Store,
    metadata: Optional[BlinkitReportStoreMetadata],
    days: list[date],
    totals_by_date: dict,
    include_vendor_rate: bool,
):
    admin_poc, contact_no = resolve_blinkit_poc(store, metadata)
    metadata = metadata or BlinkitReportStoreMetadata()
    labels = day_labels(days)
    daily = {label: int(totals_by_date.get(day, 0)) for day, label in zip(days, labels)}
    row = {
        "S.No": serial_number,
        "Region": store.region or "",
        "Entity": store.entity or "",
        "State": store.state or "",
        "city": store.city or "",
        "Citi Lead": metadata.citi_lead or "",
        "Outlet Names": store.store_name or "",
        "Outlet Id": store.id,
        "Mode of Ops": BLINKIT_MODE_OF_OPS,
        "Live Date": metadata.live_date,
        "Supply start date": metadata.supply_start_date,
        "Status": metadata.report_status or "",
        "Closed/Dropped Date": metadata.closed_dropped_date,
        "Admin POC": admin_poc,
        "Contact No": contact_no,
        "Vendor Name": BLINKIT_VENDOR_NAME,
    }
    if include_vendor_rate:
        row["Vendor Rate"] = metadata.vendor_rate
    row["Total Count"] = sum(daily.values())
    row.update(daily)
    return row


def _report_rows(db: Session, pairs, days, include_vendor_rate, serial_start=1):
    store_ids = [store.id for store, _metadata in pairs]
    totals = _monthly_totals(db, store_ids, days)
    return [
        _row(
            serial_start + index,
            store,
            metadata,
            days,
            totals.get(store.id, {}),
            include_vendor_rate,
        )
        for index, (store, metadata) in enumerate(pairs)
    ]


def build_blinkit_workbook(rows, columns) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Blinkit Distribution"
    worksheet.append(columns)
    for row in rows:
        worksheet.append([row.get(column) for column in columns])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for column_index, column in enumerate(columns, start=1):
        if column in DATE_COLUMNS:
            for row_index in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row_index, column=column_index).number_format = "dd-mmm-yyyy"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.get("/monthly", response_model=BlinkitMonthlyReportOut)
def get_monthly_blinkit_report(
    year: int = Query(..., ge=MIN_REPORT_YEAR, le=MAX_REPORT_YEAR),
    month: int = Query(..., ge=1, le=12),
    region: Optional[str] = None,
    state: Optional[str] = None,
    city: Optional[str] = None,
    entity: Optional[str] = None,
    store_id: Optional[str] = None,
    status_filter: Optional[str] = Query(None, alias="status"),
    vendor_name: Optional[str] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_access(current_user, "blinkit_reports.view", db)
    days = month_dates(year, month)
    query = _store_query(
        db,
        region=region,
        state_filter=state,
        city=city,
        entity=entity,
        store_id=store_id,
        status_filter=status_filter,
        vendor_name=vendor_name,
    )
    total = query.order_by(None).count()
    pairs = query.offset((page - 1) * page_size).limit(page_size).all()
    include_vendor_rate = _can_view_vendor_rate(current_user, db)
    rows = _report_rows(
        db,
        pairs,
        days,
        include_vendor_rate,
        serial_start=(page - 1) * page_size + 1,
    )
    return {
        "year": year,
        "month": month,
        "days": day_labels(days),
        "page": page,
        "page_size": page_size,
        "total": total,
        "rows": rows,
    }


@router.get("/monthly/export")
def export_monthly_blinkit_report(
    year: int = Query(..., ge=MIN_REPORT_YEAR, le=MAX_REPORT_YEAR),
    month: int = Query(..., ge=1, le=12),
    region: Optional[str] = None,
    state: Optional[str] = None,
    city: Optional[str] = None,
    entity: Optional[str] = None,
    store_id: Optional[str] = None,
    status_filter: Optional[str] = Query(None, alias="status"),
    vendor_name: Optional[str] = None,
    include_vendor_rate: bool = False,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_access(current_user, "blinkit_reports.export", db)
    if include_vendor_rate and not _can_view_vendor_rate(current_user, db):
        raise HTTPException(status_code=403, detail="Vendor Rate export permission required")
    days = month_dates(year, month)
    pairs = _store_query(
        db,
        region=region,
        state_filter=state,
        city=city,
        entity=entity,
        store_id=store_id,
        status_filter=status_filter,
        vendor_name=vendor_name,
    ).all()
    rows = _report_rows(db, pairs, days, include_vendor_rate)
    columns = list(BASE_COLUMNS)
    if include_vendor_rate:
        columns.append("Vendor Rate")
    columns.extend(["Total Count", *day_labels(days)])

    output = build_blinkit_workbook(rows, columns)
    filename = f"Blinkit_Distribution_{year}_{month:02d}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _text_value(value) -> Optional[str]:
    if pd.isna(value):
        return None
    text_value = str(value).strip()
    if text_value.endswith(".0") and text_value[:-2].isdigit():
        text_value = text_value[:-2]
    return text_value or None


def _date_value(value, column: str) -> Optional[date]:
    if pd.isna(value):
        return None
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        raise ValueError(f"Invalid {column}")
    return parsed.date()


def _decimal_value(value) -> Optional[Decimal]:
    if pd.isna(value):
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ValueError("Invalid Vendor Rate")


def _normalized_store_id(value) -> str:
    text_value = _text_value(value) or ""
    if text_value.isdigit():
        return text_value.lstrip("0") or "0"
    return text_value.casefold()


@router.post(
    "/metadata/import",
    response_model=BlinkitMetadataImportOut,
    status_code=status.HTTP_200_OK,
)
async def import_blinkit_metadata(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _require_access(current_user, "blinkit_reports.metadata_manage", db)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload an XLSX file")
    try:
        frame = pd.read_excel(BytesIO(await file.read()), dtype=object)
        frame.columns = [str(column).strip() for column in frame.columns]
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Unable to read workbook: {exc}")
    missing = sorted(IMPORT_COLUMNS - set(frame.columns))
    if missing:
        raise HTTPException(status_code=400, detail=f"Missing columns: {', '.join(missing)}")

    stores = db.query(Store).all()
    store_lookup = {}
    ambiguous_ids = set()
    for store in stores:
        key = _normalized_store_id(store.id)
        if key in store_lookup and store_lookup[key].id != store.id:
            ambiguous_ids.add(key)
        else:
            store_lookup[key] = store

    valid_rows = []
    errors = []
    unmatched_outlet_ids = []
    non_blinkit_outlet_ids = []
    duplicate_outlet_ids = []
    seen_outlet_keys = set()
    for frame_index, values in frame.iterrows():
        excel_row = int(frame_index) + 2
        outlet_id = _text_value(values.get("Outlet Id")) or ""
        key = _normalized_store_id(outlet_id)
        if key in seen_outlet_keys:
            duplicate_outlet_ids.append(outlet_id)
            errors.append({
                "row": excel_row,
                "outlet_id": outlet_id,
                "reason": "Duplicate Outlet Id in workbook",
            })
            continue
        seen_outlet_keys.add(key)
        store = None if key in ambiguous_ids else store_lookup.get(key)
        if not store:
            unmatched_outlet_ids.append(outlet_id)
            errors.append({"row": excel_row, "outlet_id": outlet_id, "reason": "Unknown Store ID"})
            continue
        if (store.channel or "").strip().upper() != "BLINKIT":
            non_blinkit_outlet_ids.append(outlet_id)
            errors.append({"row": excel_row, "outlet_id": outlet_id, "reason": "Store is not BLINKIT"})
            continue
        try:
            payload = {
                "citi_lead": _text_value(values.get("Citi Lead")),
                "live_date": _date_value(values.get("Live Date"), "Live Date"),
                "supply_start_date": _date_value(values.get("Supply start date"), "Supply start date"),
                "report_status": _text_value(values.get("Status")),
                "vendor_rate": _decimal_value(values.get("Vendor Rate")),
            }
        except ValueError as exc:
            errors.append({"row": excel_row, "outlet_id": outlet_id, "reason": str(exc)})
            continue
        valid_rows.append((store, payload))

    matched_ids = {store.id for store, _payload in valid_rows}
    existing = {
        item.store_id: item
        for item in db.query(BlinkitReportStoreMetadata)
        .filter(BlinkitReportStoreMetadata.store_id.in_(matched_ids))
        .all()
    } if matched_ids else {}
    matched_blinkit_outlet_ids = sorted(matched_ids)
    rows_to_create = sum(store.id not in existing for store, _payload in valid_rows)
    rows_to_update = len(valid_rows) - rows_to_create
    created = 0
    updated = 0
    try:
        for store, payload in valid_rows:
            metadata = existing.get(store.id)
            if metadata is None:
                metadata = BlinkitReportStoreMetadata(
                    store_id=store.id,
                    created_by_id=current_user.id,
                )
                db.add(metadata)
                existing[store.id] = metadata
                created += 1
            else:
                updated += 1
            for field, value in payload.items():
                setattr(metadata, field, value)
            metadata.updated_by_id = current_user.id
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Metadata import failed: {exc}")

    return {
        "total_rows": len(frame),
        "matched_blinkit_outlet_ids": matched_blinkit_outlet_ids,
        "unmatched_outlet_ids": list(dict.fromkeys(unmatched_outlet_ids)),
        "non_blinkit_outlet_ids": list(dict.fromkeys(non_blinkit_outlet_ids)),
        "duplicate_outlet_ids": list(dict.fromkeys(duplicate_outlet_ids)),
        "rows_to_create": rows_to_create,
        "rows_to_update": rows_to_update,
        "invalid_rows": len(errors),
        "created": created,
        "updated": updated,
        "skipped": len(errors),
        "errors": errors,
    }
