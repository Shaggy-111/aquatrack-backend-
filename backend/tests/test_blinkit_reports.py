import asyncio
import ast
import inspect
import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import backend.models
from backend.database.db import Base
from backend.models.blinkit_report_store_metadata import BlinkitReportStoreMetadata
from backend.models.order import Order, OrderStatusEnum
from backend.models.partner_store import PartnerStore
from backend.models.permission import Permission
from backend.models.store import Store
from backend.models.user import RoleEnum, User
from backend.routes import blinkit_reports
from backend.routes.employee_menu import get_my_sidebar
from backend.utils.permission_seeder import (
    FINALIZED_PERMISSION_MATRIX,
    build_permission_rows,
    build_sidebar_modules,
    seed_default_permissions,
)


class BlinkitReportTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.admin = User(
            email="blinkit-admin@test.local",
            password="test",
            full_name="Admin",
            role=RoleEnum.superadmin,
            status="active",
        )
        self.db.add(self.admin)
        self.db.flush()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    def add_store(self, store_id, *, channel="BLINKIT", region="North", state="Delhi", city="Delhi"):
        store = Store(
            id=store_id,
            store_name=f"Store {store_id}",
            channel=channel,
            region=region,
            state=state,
            city=city,
            entity="Blinkit Entity",
            is_active=True,
        )
        self.db.add(store)
        self.db.flush()
        return store

    def add_order(self, store, *, status, delivered, confirmed, when, source):
        order = Order(
            partner_id=self.admin.id,
            store_id=store.id,
            order_details=str(delivered),
            total_amount=0,
            status=status,
            order_date=when,
            bottles_delivered=delivered,
            confirmed_bottles=confirmed,
            ordered_by_type=source,
        )
        self.db.add(order)
        self.db.flush()
        return order

    def add_mapped_user(self, store, user_id, *, role=RoleEnum.partner, name=None, mobile=None):
        user = User(
            id=user_id,
            email=f"poc-{user_id}@test.local",
            password="test",
            full_name=name or f"POC {user_id}",
            mobile_number=mobile or f"900000{user_id:04d}",
            role=role,
            status="active",
        )
        self.db.add(user)
        self.db.flush()
        self.db.add(PartnerStore(user_id=user.id, store_id=store.id))
        self.db.flush()
        self.db.expire(store, ["partner_associations"])
        return user

    def test_month_lengths_and_ist_boundaries(self):
        self.assertEqual(len(blinkit_reports.month_dates(2027, 2)), 28)
        self.assertEqual(len(blinkit_reports.month_dates(2028, 2)), 29)
        self.assertEqual(len(blinkit_reports.month_dates(2026, 4)), 30)
        self.assertEqual(len(blinkit_reports.month_dates(2026, 7)), 31)
        self.assertEqual(blinkit_reports.day_labels(blinkit_reports.month_dates(2026, 7))[-1], "31-Jul")

        start, end = blinkit_reports.ist_date_range_utc(date(2026, 12, 1), date(2026, 12, 31))
        self.assertEqual(start, datetime(2026, 11, 30, 18, 30, tzinfo=timezone.utc))
        self.assertEqual(end, datetime(2026, 12, 31, 18, 30, tzinfo=timezone.utc))

    def test_year_and_month_validation_is_bounded(self):
        self.assertEqual(blinkit_reports.month_dates(2020, 1)[0], date(2020, 1, 1))
        self.assertEqual(blinkit_reports.month_dates(2100, 12)[-1], date(2100, 12, 31))
        for invalid_year in (0, 2019, 2101, 10000):
            with self.subTest(year=invalid_year), self.assertRaises(HTTPException) as raised:
                blinkit_reports.month_dates(invalid_year, 1)
            self.assertEqual(raised.exception.status_code, 400)
        for invalid_month in (0, 13):
            with self.subTest(month=invalid_month), self.assertRaises(HTTPException) as raised:
                blinkit_reports.month_dates(2026, invalid_month)
            self.assertEqual(raised.exception.status_code, 400)

    def test_quantities_sources_statuses_and_same_day_aggregation(self):
        store = self.add_store("BL-1")
        included = [
            self.add_order(store, status=OrderStatusEnum.delivered_confirmed, delivered=10, confirmed=7,
                           when=datetime(2026, 7, 5, 5, 0, tzinfo=timezone.utc), source="VIRTUAL_CARD"),
            self.add_order(store, status=OrderStatusEnum.delivered_confirmed, delivered=9, confirmed=0,
                           when=datetime(2026, 7, 5, 6, 0, tzinfo=timezone.utc), source="QR_DELIVERY"),
            self.add_order(store, status=OrderStatusEnum.delivered_pending_confirmation, delivered=4, confirmed=99,
                           when=datetime(2026, 7, 5, 7, 0, tzinfo=timezone.utc), source="ADMIN"),
            self.add_order(store, status=OrderStatusEnum.delivered_confirmed, delivered=3, confirmed=3,
                           when=datetime(2026, 7, 5, 8, 0, tzinfo=timezone.utc), source="EXCEL_UPLOAD"),
        ]
        self.add_order(store, status=OrderStatusEnum.cancelled, delivered=100, confirmed=100,
                       when=datetime(2026, 7, 5, 9, 0, tzinfo=timezone.utc), source="ADMIN")
        self.db.commit()

        totals = blinkit_reports._monthly_totals(
            self.db, [store.id], blinkit_reports.month_dates(2026, 7)
        )
        self.assertEqual(totals[store.id], {date(2026, 7, 5): 14})
        self.assertEqual(len(included), 4)  # Orders are aggregated, never deduplicated.

        row = blinkit_reports._row(
            1, store, None, blinkit_reports.month_dates(2026, 7), totals[store.id], False
        )
        self.assertEqual(row["5-Jul"], 14)
        self.assertEqual(row["6-Jul"], 0)
        self.assertEqual(row["Total Count"], sum(row[label] for label in blinkit_reports.day_labels(blinkit_reports.month_dates(2026, 7))))
        self.assertNotIn("Vendor Rate", row)
        self.assertNotIn("Duplicacy Check", row)
        self.assertNotIn("daily_counts", row)
        self.assertIsInstance(row["5-Jul"], int)

    def test_midnight_final_day_and_next_month_are_grouped_correctly(self):
        store = self.add_store("BL-TIME")
        self.add_order(store, status=OrderStatusEnum.delivered_confirmed, delivered=2, confirmed=2,
                       when=datetime(2026, 6, 30, 18, 30, tzinfo=timezone.utc), source="ADMIN")  # 1-Jul 00:00 IST
        self.add_order(store, status=OrderStatusEnum.delivered_confirmed, delivered=3, confirmed=3,
                       when=datetime(2026, 7, 31, 18, 29, tzinfo=timezone.utc), source="ADMIN")  # 31-Jul 23:59 IST
        self.add_order(store, status=OrderStatusEnum.delivered_confirmed, delivered=99, confirmed=99,
                       when=datetime(2026, 7, 31, 18, 30, tzinfo=timezone.utc), source="ADMIN")  # 1-Aug
        self.db.commit()
        totals = blinkit_reports._monthly_totals(self.db, [store.id], blinkit_reports.month_dates(2026, 7))[store.id]
        self.assertEqual(totals, {date(2026, 7, 1): 2, date(2026, 7, 31): 3})

    def test_only_blinkit_filters_and_deterministic_order(self):
        self.add_store("3", region="West", state="B", city="C")
        first = self.add_store("1", region="North", state="A", city="B")
        second = self.add_store("2", region="North", state="A", city="A")
        self.add_store("NON", channel="ZEPTO")
        self.db.add_all([
            BlinkitReportStoreMetadata(store_id=first.id, report_status="Live", vendor_name="Vendor A"),
            BlinkitReportStoreMetadata(store_id=second.id, report_status="Closed", vendor_name="Vendor B"),
        ])
        self.db.commit()
        pairs = blinkit_reports._store_query(
            self.db, region=None, state_filter=None, city=None, entity=None,
            store_id=None, status_filter=None, vendor_name=None,
        ).all()
        self.assertEqual([store.id for store, _ in pairs], ["2", "1", "3"])
        filtered = blinkit_reports._store_query(
            self.db, region="north", state_filter="a", city="b", entity="blinkit entity",
            store_id="1", status_filter="live",
            vendor_name="VEE KAY AQUATECH Private Limited",
        ).all()
        self.assertEqual([store.id for store, _ in filtered], ["1"])
        mismatched_vendor = blinkit_reports._store_query(
            self.db, region=None, state_filter=None, city=None, entity=None,
            store_id=None, status_filter=None, vendor_name="Vendor A",
        ).all()
        self.assertEqual(mismatched_vendor, [])

    def test_fixed_mode_and_vendor_override_metadata_in_json_and_xlsx(self):
        store = self.add_store("BL-FIXED-PRESENTATION")
        metadata = BlinkitReportStoreMetadata(
            store_id=store.id,
            mode_of_ops="Metadata Mode",
            vendor_name="Metadata Vendor",
            vendor_rate=Decimal("12.3400"),
        )
        original = (metadata.mode_of_ops, metadata.vendor_name, metadata.vendor_rate)
        row = blinkit_reports._row(
            1,
            store,
            metadata,
            [date(2026, 7, 1), date(2026, 7, 2)],
            {date(2026, 7, 1): 9},
            True,
        )
        self.assertEqual(row["Mode of Ops"], "Real Estate")
        self.assertEqual(
            row["Vendor Name"], "VEE KAY AQUATECH Private Limited"
        )
        self.assertEqual(row["Vendor Rate"], Decimal("12.3400"))
        self.assertEqual((row["1-Jul"], row["2-Jul"], row["Total Count"]), (9, 0, 9))

        columns = [
            *blinkit_reports.BASE_COLUMNS,
            "Vendor Rate",
            "Total Count",
            "1-Jul",
            "2-Jul",
        ]
        workbook = load_workbook(
            blinkit_reports.build_blinkit_workbook([row], columns)
        )
        headers = [cell.value for cell in workbook.active[1]]
        exported = {
            headers[index]: workbook.active.cell(2, index + 1).value
            for index in range(len(headers))
        }
        self.assertEqual(exported["Mode of Ops"], "Real Estate")
        self.assertEqual(
            exported["Vendor Name"], "VEE KAY AQUATECH Private Limited"
        )
        self.assertEqual(exported["Vendor Rate"], 12.34)
        self.assertEqual(
            (metadata.mode_of_ops, metadata.vendor_name, metadata.vendor_rate), original
        )

    def test_poc_metadata_overrides_and_blank_values_fall_back_to_same_partner(self):
        store = self.add_store("BL-POC")
        partner = self.add_mapped_user(
            store, 101, name="Mapped Partner", mobile="9929202526"
        )

        overrides = BlinkitReportStoreMetadata(
            admin_poc="Custom Admin",
            contact_no="9999999999",
        )
        self.assertEqual(
            blinkit_reports.resolve_blinkit_poc(store, overrides),
            ("Custom Admin", "9999999999"),
        )

        blank = BlinkitReportStoreMetadata(admin_poc="  ", contact_no="")
        self.assertEqual(
            blinkit_reports.resolve_blinkit_poc(store, blank),
            (partner.full_name, partner.mobile_number),
        )

        partial = BlinkitReportStoreMetadata(
            admin_poc="Custom Admin",
            contact_no=" ",
        )
        self.assertEqual(
            blinkit_reports.resolve_blinkit_poc(store, partial),
            ("Custom Admin", partner.mobile_number),
        )
        self.assertEqual(
            blinkit_reports.resolve_blinkit_poc(store, None),
            (partner.full_name, partner.mobile_number),
        )
        self.assertEqual((blank.admin_poc, blank.contact_no), ("  ", ""))
        self.assertEqual((partial.admin_poc, partial.contact_no), ("Custom Admin", " "))

    def test_poc_selection_prefers_partner_then_lowest_user_id_deterministically(self):
        store = self.add_store("BL-MULTI")
        self.add_mapped_user(
            store, 10, role=RoleEnum.employee, name="Non Partner", mobile="100"
        )
        selected = self.add_mapped_user(
            store, 30, role=RoleEnum.partner, name="Selected Partner", mobile="300"
        )
        self.add_mapped_user(
            store, 40, role=RoleEnum.partner, name="Later Partner", mobile="400"
        )
        self.assertEqual(
            blinkit_reports.resolve_blinkit_poc(store, None),
            (selected.full_name, selected.mobile_number),
        )

        unmapped = self.add_store("BL-NO-POC")
        self.assertEqual(blinkit_reports.resolve_blinkit_poc(unmapped, None), ("", ""))

    def test_json_and_xlsx_use_identical_resolved_poc_values(self):
        store = self.add_store("BL-POC-XLSX")
        self.add_mapped_user(store, 111, name="Excel POC", mobile="9111111111")
        row = blinkit_reports._row(
            1, store, None, [date(2026, 7, 1)], {}, False
        )
        columns = [*blinkit_reports.BASE_COLUMNS, "Total Count", "1-Jul"]
        workbook = load_workbook(
            blinkit_reports.build_blinkit_workbook([row], columns)
        )
        headers = [cell.value for cell in workbook.active[1]]
        exported = {
            headers[index]: workbook.active.cell(2, index + 1).value
            for index in range(len(headers))
        }
        self.assertEqual(exported["Admin POC"], row["Admin POC"])
        self.assertEqual(exported["Contact No"], row["Contact No"])

    def test_vendor_rate_security_and_json_omission(self):
        employee = SimpleNamespace(id=20, role=RoleEnum.employee)
        with patch.object(blinkit_reports, "get_employee_permissions", return_value=[]):
            with self.assertRaises(HTTPException):
                blinkit_reports._require_access(employee, "blinkit_reports.view", self.db)
            self.assertFalse(blinkit_reports._can_view_vendor_rate(employee, self.db))
        with patch.object(
            blinkit_reports,
            "get_employee_permissions",
            return_value=["blinkit_reports.view", "blinkit_reports.vendor_rate_export"],
        ):
            blinkit_reports._require_access(employee, "blinkit_reports.view", self.db)
            self.assertTrue(blinkit_reports._can_view_vendor_rate(employee, self.db))
        with self.assertRaises(HTTPException):
            blinkit_reports._require_access(SimpleNamespace(role=RoleEnum.partner), "blinkit_reports.view", self.db)

        store = self.add_store("BL-RATE")
        metadata = BlinkitReportStoreMetadata(store_id=store.id, vendor_rate=Decimal("12.3400"))
        hidden = blinkit_reports._row(1, store, metadata, [date(2026, 7, 1)], {}, False)
        shown = blinkit_reports._row(1, store, metadata, [date(2026, 7, 1)], {}, True)
        self.assertNotIn("Vendor Rate", hidden)
        self.assertEqual(shown["Vendor Rate"], Decimal("12.3400"))

    def test_blinkit_wildcards_grant_all_blinkit_operations(self):
        employee = SimpleNamespace(id=21, role=RoleEnum.employee)
        required = (
            "blinkit_reports.view",
            "blinkit_reports.export",
            "blinkit_reports.metadata_manage",
        )
        for wildcard in ("*", "*.*"):
            with self.subTest(wildcard=wildcard), patch.object(
                blinkit_reports, "get_employee_permissions", return_value=[wildcard]
            ):
                for permission_key in required:
                    blinkit_reports._require_access(employee, permission_key, self.db)
                self.assertTrue(blinkit_reports._can_view_vendor_rate(employee, self.db))

                unknown_row = {column: None for column in blinkit_reports.IMPORT_COLUMNS}
                unknown_row["Outlet Id"] = "UNKNOWN"
                result = asyncio.run(
                    blinkit_reports.import_blinkit_metadata(
                        file=self.workbook_upload([unknown_row]),
                        db=self.db,
                        current_user=employee,
                    )
                )
                self.assertEqual(result["skipped"], 1)

    def test_superadmin_and_other_role_authorization_unchanged(self):
        blinkit_reports._require_access(self.admin, "blinkit_reports.view", self.db)
        self.assertTrue(blinkit_reports._can_view_vendor_rate(self.admin, self.db))
        for role in (RoleEnum.partner, RoleEnum.delivery_manager):
            with self.subTest(role=role), self.assertRaises(HTTPException):
                blinkit_reports._require_access(SimpleNamespace(role=role), "blinkit_reports.view", self.db)

    def test_export_columns_types_and_dates(self):
        store = self.add_store("BL-XLSX")
        metadata = BlinkitReportStoreMetadata(
            store_id=store.id,
            live_date=date(2026, 1, 2),
            vendor_rate=Decimal("9.5000"),
        )
        days = blinkit_reports.month_dates(2026, 2)
        row = blinkit_reports._row(1, store, metadata, days, {date(2026, 2, 1): 5}, True)
        columns = [*blinkit_reports.BASE_COLUMNS, "Vendor Rate", "Total Count", *blinkit_reports.day_labels(days)]
        workbook = load_workbook(blinkit_reports.build_blinkit_workbook([row], columns))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, columns)
        self.assertNotIn("Duplicacy Check", headers)
        self.assertEqual(headers.index("Vendor Rate"), headers.index("Vendor Name") + 1)
        self.assertEqual(headers.index("Total Count"), headers.index("Vendor Rate") + 1)
        self.assertEqual(len(headers[-28:]), 28)
        values = {headers[index]: sheet.cell(2, index + 1).value for index in range(len(headers))}
        self.assertIsInstance(values["Live Date"], (date, datetime))
        self.assertIsInstance(values["Total Count"], int)
        self.assertIsInstance(values["1-Feb"], int)

    def workbook_upload(self, rows):
        output = BytesIO()
        pd.DataFrame(rows).to_excel(output, index=False)
        output.seek(0)
        return UploadFile(filename="metadata.xlsx", file=output)

    def test_metadata_import_accepts_only_xlsx(self):
        valid = {column: None for column in blinkit_reports.IMPORT_COLUMNS}
        valid["Outlet Id"] = "UNKNOWN"
        accepted = asyncio.run(
            blinkit_reports.import_blinkit_metadata(
                file=self.workbook_upload([valid]), db=self.db, current_user=self.admin
            )
        )
        self.assertEqual(accepted["skipped"], 1)

        for filename in ("metadata.xls", "metadata.csv", "metadata.txt"):
            with self.subTest(filename=filename), self.assertRaises(HTTPException) as raised:
                asyncio.run(
                    blinkit_reports.import_blinkit_metadata(
                        file=UploadFile(filename=filename, file=BytesIO(b"unsupported")),
                        db=self.db,
                        current_user=self.admin,
                    )
                )
            self.assertEqual(raised.exception.status_code, 400)

    def test_metadata_import_rejects_unknown_non_blinkit_and_upserts_decimal(self):
        blinkit = self.add_store("001")
        self.add_store("Z-1", channel="ZEPTO")
        row = {
            "Outlet Id": "1", "Citi Lead": "Lead", "Mode of Ops": "Direct",
            "Live Date": date(2026, 1, 1), "Supply start date": date(2026, 1, 2),
            "Status": "Live", "Closed/Dropped Date": None, "Admin POC": "Admin",
            "Contact No": "09999999999", "Vendor Name": "Vendor", "Vendor Rate": "12.3456",
        }
        result = asyncio.run(blinkit_reports.import_blinkit_metadata(
            file=self.workbook_upload([row, {**row, "Outlet Id": "UNKNOWN"}, {**row, "Outlet Id": "Z-1"}]),
            db=self.db, current_user=self.admin,
        ))
        self.assertEqual((result["created"], result["updated"], result["skipped"]), (1, 0, 2))
        metadata = self.db.query(BlinkitReportStoreMetadata).filter_by(store_id=blinkit.id).one()
        self.assertEqual(metadata.vendor_rate, Decimal("12.3456"))

        row["Vendor Rate"] = "7.2500"
        second = asyncio.run(blinkit_reports.import_blinkit_metadata(
            file=self.workbook_upload([row]), db=self.db, current_user=self.admin,
        ))
        self.assertEqual((second["created"], second["updated"]), (0, 1))
        self.assertEqual(self.db.query(BlinkitReportStoreMetadata).count(), 1)
        self.assertEqual(self.db.query(BlinkitReportStoreMetadata).one().vendor_rate, Decimal("7.2500"))

    def test_metadata_import_ignores_outlet_names_and_updates_only_allowed_fields(self):
        store = self.add_store("BL-IMPORT-ONLY")
        original_store_name = store.store_name
        existing = BlinkitReportStoreMetadata(
            store_id=store.id,
            citi_lead="Old Lead",
            mode_of_ops="Legacy Mode",
            live_date=date(2025, 1, 1),
            supply_start_date=date(2025, 1, 2),
            report_status="Old",
            closed_dropped_date=date(2025, 1, 3),
            admin_poc="Custom POC",
            contact_no="9999999999",
            vendor_name="Legacy Vendor",
            vendor_rate=Decimal("1.0000"),
        )
        self.db.add(existing)
        self.db.commit()

        valid_row = {
            "Outlet Id": store.id,
            "Outlet Names": "DELIBERATELY INCORRECT NAME",
            "Citi Lead": "New Lead",
            "Live Date": date(2026, 7, 1),
            "Supply start date": date(2026, 7, 2),
            "Status": "Live",
            "Vendor Rate": "8.7500",
            "Admin POC": "Must Be Ignored",
            "Contact No": "Must Be Ignored",
            "Vendor Name": "Must Be Ignored",
            "Mode of Ops": "Must Be Ignored",
        }
        name_only_match = {
            **valid_row,
            "Outlet Id": "UNKNOWN-BUT-SAME-NAME",
            "Outlet Names": original_store_name,
        }
        result = asyncio.run(
            blinkit_reports.import_blinkit_metadata(
                file=self.workbook_upload([valid_row, name_only_match]),
                db=self.db,
                current_user=self.admin,
            )
        )
        self.assertEqual(result["matched_blinkit_outlet_ids"], [store.id])
        self.assertEqual(result["unmatched_outlet_ids"], ["UNKNOWN-BUT-SAME-NAME"])
        self.assertEqual((result["rows_to_create"], result["rows_to_update"]), (0, 1))

        self.db.refresh(store)
        self.db.refresh(existing)
        self.assertEqual(store.store_name, original_store_name)
        self.assertEqual(existing.citi_lead, "New Lead")
        self.assertEqual(existing.live_date, date(2026, 7, 1))
        self.assertEqual(existing.supply_start_date, date(2026, 7, 2))
        self.assertEqual(existing.report_status, "Live")
        self.assertEqual(existing.vendor_rate, Decimal("8.7500"))
        self.assertEqual(existing.mode_of_ops, "Legacy Mode")
        self.assertEqual(existing.closed_dropped_date, date(2025, 1, 3))
        self.assertEqual(existing.admin_poc, "Custom POC")
        self.assertEqual(existing.contact_no, "9999999999")
        self.assertEqual(existing.vendor_name, "Legacy Vendor")

        repeated = asyncio.run(
            blinkit_reports.import_blinkit_metadata(
                file=self.workbook_upload([valid_row]),
                db=self.db,
                current_user=self.admin,
            )
        )
        self.assertEqual((repeated["created"], repeated["updated"]), (0, 1))
        self.assertEqual(
            self.db.query(BlinkitReportStoreMetadata)
            .filter(BlinkitReportStoreMetadata.store_id == store.id)
            .count(),
            1,
        )

    def test_metadata_import_reports_duplicate_outlet_ids(self):
        store = self.add_store("BL-DUPLICATE")
        row = {
            "Outlet Id": store.id,
            "Outlet Names": "Ignored",
            "Citi Lead": "Lead",
            "Live Date": None,
            "Supply start date": None,
            "Status": "Live",
            "Vendor Rate": "2.5000",
        }
        result = asyncio.run(
            blinkit_reports.import_blinkit_metadata(
                file=self.workbook_upload([row, {**row, "Citi Lead": "Second"}]),
                db=self.db,
                current_user=self.admin,
            )
        )
        self.assertEqual(result["duplicate_outlet_ids"], [store.id])
        self.assertEqual(result["invalid_rows"], 1)
        self.assertEqual((result["created"], result["updated"], result["skipped"]), (1, 0, 1))
        self.assertEqual(
            self.db.query(BlinkitReportStoreMetadata).filter_by(store_id=store.id).one().citi_lead,
            "Lead",
        )

    def test_permissions_sidebar_idempotency_and_existing_report_isolation(self):
        rows = {row["permission_key"]: row for row in build_permission_rows()}
        self.assertEqual(rows["blinkit_reports.vendor_rate_export"]["display_name"], "Export Blinkit Vendor Rate")
        sidebar = build_sidebar_modules()
        blinkit = [item for item in sidebar if item["module"] == "Blinkit Report"]
        self.assertEqual(len(blinkit), 1)
        self.assertEqual(
            {key: blinkit[0][key] for key in ("module", "permission", "icon", "tab", "route", "group", "order")},
            {"module": "Blinkit Report", "permission": "blinkit_reports.view", "icon": "assessment",
             "tab": "blinkitReports", "route": "blinkitReports", "group": "Reports", "order": 322},
        )
        labels = [item["module"] for item in sidebar if item["group"] == "Reports"]
        self.assertLess(labels.index("Delivery Reports"), labels.index("Blinkit Report"))
        self.assertLess(labels.index("Blinkit Report"), labels.index("Monthly Virtual Cards"))
        delivery = next(item for item in sidebar if item["module"] == "Delivery Reports")
        self.assertEqual((delivery["permission"], delivery["tab"], delivery["order"]),
                         ("delivery_reports.view", "deliveryReports", 320))
        seed_default_permissions(self.db)
        seed_default_permissions(self.db)
        self.assertEqual(self.db.query(Permission).filter(Permission.permission_key == "blinkit_reports.view").count(), 1)
        self.assertNotIn("db.add", inspect.getsource(blinkit_reports.get_monthly_blinkit_report))
        self.assertNotIn("Report(", inspect.getsource(blinkit_reports.get_monthly_blinkit_report))

    def test_blinkit_permission_sort_orders_do_not_shift_existing_catalog(self):
        actual = {row["permission_key"]: row["sort_order"] for row in build_permission_rows()}
        expected_sort_order = 10
        for module in FINALIZED_PERMISSION_MATRIX:
            module_key = module.get("key")
            if not module_key:
                import re
                module_key = re.sub(r"[^a-z0-9]+", "_", module["module"].lower()).strip("_")
            if module_key == "blinkit_reports":
                continue
            for action in module["actions"]:
                action_key = str(action).lower().replace(" ", "_")
                self.assertEqual(actual[f"{module_key}.{action_key}"], expected_sort_order)
                expected_sort_order += 10
        self.assertEqual(
            [actual[f"blinkit_reports.{action}"] for action in
             ("view", "export", "metadata_manage", "vendor_rate_export")],
            [900010, 900020, 900030, 900040],
        )

    def test_report_rows_bulk_load_orders_once(self):
        stores = [self.add_store(f"BL-PERF-{index}") for index in range(3)]
        statement_count = 0

        def count_statement(*_args):
            nonlocal statement_count
            statement_count += 1

        event.listen(self.engine, "before_cursor_execute", count_statement)
        try:
            pairs = blinkit_reports._store_query(
                self.db,
                region=None,
                state_filter=None,
                city=None,
                entity=None,
                store_id=None,
                status_filter=None,
                vendor_name=None,
            ).all()
            rows = blinkit_reports._report_rows(
                self.db,
                pairs,
                blinkit_reports.month_dates(2026, 7),
                False,
            )
        finally:
            event.remove(self.engine, "before_cursor_execute", count_statement)
        self.assertEqual(len(rows), 3)
        self.assertEqual(statement_count, 2)

    def test_poc_relationships_are_bulk_loaded_without_n_plus_one(self):
        stores = [self.add_store(f"BL-POC-PERF-{index}") for index in range(3)]
        for index, store in enumerate(stores, start=201):
            self.add_mapped_user(store, index)
        self.db.commit()
        statement_count = 0

        def count_statement(*_args):
            nonlocal statement_count
            statement_count += 1

        event.listen(self.engine, "before_cursor_execute", count_statement)
        try:
            pairs = blinkit_reports._store_query(
                self.db,
                region=None,
                state_filter=None,
                city=None,
                entity=None,
                store_id=None,
                status_filter=None,
                vendor_name=None,
            ).all()
            resolved = [
                blinkit_reports.resolve_blinkit_poc(store, metadata)
                for store, metadata in pairs
                if store.id.startswith("BL-POC-PERF-")
            ]
        finally:
            event.remove(self.engine, "before_cursor_execute", count_statement)
        self.assertEqual(len(resolved), 3)
        self.assertEqual(statement_count, 1)

    def test_migration_extends_current_head_and_chain_stays_single_head(self):
        revisions = {}
        version_dir = Path(__file__).parents[1] / "alembic" / "versions"
        for path in version_dir.glob("*.py"):
            tree = ast.parse(path.read_text(encoding="utf-8-sig"))
            values = {}
            for node in tree.body:
                if isinstance(node, (ast.Assign, ast.AnnAssign)):
                    targets = node.targets if isinstance(node, ast.Assign) else [node.target]
                    for target in targets:
                        if isinstance(target, ast.Name) and target.id in {"revision", "down_revision"}:
                            try:
                                values[target.id] = ast.literal_eval(node.value)
                            except (ValueError, TypeError):
                                pass
            if values.get("revision"):
                revisions[values["revision"]] = values.get("down_revision")
        self.assertEqual(revisions["d8b3f6a2c9e1"], "c3e7a9b1d5f2")
        referenced = {
            parent
            for parent_value in revisions.values()
            for parent in (parent_value if isinstance(parent_value, (tuple, list)) else [parent_value])
            if parent
        }
        self.assertEqual(set(revisions) - referenced, {"d8b3f6a2c9e1"})


if __name__ == "__main__":
    unittest.main()
