import io
import unittest
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

from fastapi import HTTPException, UploadFile
from PIL import Image
import openpyxl
from sqlalchemy import create_engine, event
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import backend.models
from backend.database.db import Base
from backend.models.employee_permission import EmployeePermission
from backend.models.employee_store_mapping import EmployeeStoreMapping
from backend.models.permission import Permission
from backend.models.store import Store
from backend.models.store_compliance_monthly import (
    StoreComplianceAttachment,
    StoreComplianceAttachmentType,
    StoreComplianceInvoice,
    StoreCompliancePayment,
    StoreCompliancePaymentStatus,
    StoreMonthlyCardStatus,
    StoreMonthlyCompliance,
)
from backend.models.user import RoleEnum, User
from backend.routes.store_compliance_monthly import (
    create_compliance_invoice_endpoint,
    create_monthly_compliance_endpoint,
    export_monthly_compliance_excel_endpoint,
    get_monthly_compliance_detail_endpoint,
    get_monthly_compliance_summary_endpoint,
    list_monthly_compliances_endpoint,
    record_compliance_payment_endpoint,
    update_compliance_invoice_endpoint,
    upload_compliance_attachment_endpoint,
    upload_monthly_card_endpoint,
    void_compliance_attachment_endpoint,
)
from backend.schemas.store_compliance_monthly import (
    StoreComplianceAttachmentVoidRequest,
    StoreComplianceInvoiceCreate,
    StoreComplianceInvoiceUpdate,
    StoreCompliancePaymentCreate,
    StoreMonthlyComplianceCreate,
)


def _generate_test_image(color="blue", size=(100, 100)) -> bytes:
    buf = io.BytesIO()
    img = Image.new("RGB", size, color=color)
    img.save(buf, format="JPEG")
    return buf.getvalue()


class StoreComplianceMonthlyModuleTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )

        @event.listens_for(self.engine, "connect")
        def set_sqlite_pragma(dbapi_connection, connection_record):
            cursor = dbapi_connection.cursor()
            cursor.execute("PRAGMA foreign_keys=ON")
            cursor.close()

        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()

        # 1. SuperAdmin User
        self.superadmin = User(
            id=1,
            email="superadmin@aquatrack.local",
            password="test",
            full_name="Global SuperAdmin",
            role=RoleEnum.superadmin,
            status="active",
        )

        # 2. Employee User (Scoped to Store 1)
        self.employee = User(
            id=2,
            email="emp@aquatrack.local",
            password="test",
            full_name="Operations Employee",
            role=RoleEnum.employee,
            status="active",
        )

        # 3. Other Employee (Unassigned)
        self.other_employee = User(
            id=3,
            email="other_emp@aquatrack.local",
            password="test",
            full_name="Other Employee",
            role=RoleEnum.employee,
            status="active",
        )

        # 4. Partner User
        self.partner = User(
            id=4,
            email="partner@aquatrack.local",
            password="test",
            full_name="Store Partner",
            role=RoleEnum.partner,
            status="active",
        )

        # 5. Delivery Manager User
        self.dm = User(
            id=5,
            email="dm@aquatrack.local",
            password="test",
            full_name="Delivery Manager North",
            role=RoleEnum.delivery_manager,
            status="active",
        )

        self.db.add_all([
            self.superadmin,
            self.employee,
            self.other_employee,
            self.partner,
            self.dm,
        ])
        self.db.flush()

        # Seed Stores
        self.store_1 = Store(
            id="STR-GGN-01",
            store_name="Gurgaon Sector 29",
            channel="BLINKIT",
            region="NORTH",
            entity="BCPL",
            city="Gurgaon",
            state="Haryana",
            assigned_manager_id=self.dm.id,
            is_active=True,
        )
        self.store_2 = Store(
            id="STR-BLR-02",
            store_name="Bangalore Indiranagar",
            channel="ZEPTO",
            region="SOUTH",
            entity="ZEPTO",
            city="Bangalore",
            state="Karnataka",
            is_active=True,
        )
        self.db.add_all([self.store_1, self.store_2])
        self.db.flush()

        # Assign Store 1 to Employee (Store 2 unassigned)
        self.db.add(
            EmployeeStoreMapping(
                employee_id=self.employee.id,
                store_id=self.store_1.id,
                assigned_by=self.superadmin.id,
            )
        )

        # Seed Permissions
        perm_keys = [
            "store_compliance.view",
            "store_compliance.create",
            "store_compliance.edit",
            "store_compliance.upload",
            "store_compliance.invoice_manage",
            "store_compliance.record_payment",
            "store_compliance.export",
        ]
        self.perms = {}
        for idx, key in enumerate(perm_keys, start=1):
            p = Permission(
                id=idx,
                module="Store Compliance",
                action=key.split(".")[1],
                permission_key=key,
                display_name=key,
                is_active=True,
            )
            self.db.add(p)
            self.perms[key] = p
        self.db.flush()

        # Assign employee permissions for view, create, upload, export
        for key in ["store_compliance.view", "store_compliance.create", "store_compliance.upload", "store_compliance.export"]:
            self.db.add(
                EmployeePermission(
                    employee_id=self.employee.id,
                    permission_id=self.perms[key].id,
                    assigned_by=self.superadmin.id,
                )
            )

        self.db.commit()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    # =========================================================================
    # 1. NON-DESTRUCTIVE FK & NO-CASCADE TESTS
    # =========================================================================

    async def test_parent_monthly_compliance_cannot_cascade_delete_invoice(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-DEL-1",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("2000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        # Verify that deleting monthly compliance directly fails with IntegrityError (RESTRICT FK)
        c_db = self.db.query(StoreMonthlyCompliance).filter(StoreMonthlyCompliance.id == c.id).first()
        with self.assertRaises(IntegrityError):
            with self.db.begin_nested():
                self.db.delete(c_db)
                self.db.flush()

    async def test_invoice_cannot_cascade_delete_payments(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        inv = await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-DEL-2",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("5000.00"),
            db=self.db,
            current_user=self.superadmin,
        )
        record_compliance_payment_endpoint(
            compliance_id=c.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("1000.00"), payment_date=datetime.now(timezone.utc)),
            db=self.db,
            current_user=self.superadmin,
        )

        inv_db = self.db.query(StoreComplianceInvoice).filter(StoreComplianceInvoice.id == inv.id).first()
        # Verify that deleting invoice directly fails with IntegrityError (RESTRICT FK)
        with self.assertRaises(IntegrityError):
            with self.db.begin_nested():
                self.db.delete(inv_db)
                self.db.flush()

    # =========================================================================
    # 2. COMPLIANCE CREATION TESTS
    # =========================================================================

    def test_create_monthly_compliance_success(self):
        payload = StoreMonthlyComplianceCreate(
            store_id="STR-GGN-01",
            year=2026,
            month=10,
            card_remarks="Initial October card pending",
        )
        res = create_monthly_compliance_endpoint(payload=payload, db=self.db, current_user=self.superadmin)
        self.assertEqual(res.store_id, "STR-GGN-01")
        self.assertEqual(res.year, 2026)
        self.assertEqual(res.month, 10)
        self.assertEqual(res.card_status, "PENDING_CARD")
        self.assertEqual(res.store_name, "Gurgaon Sector 29")
        self.assertEqual(res.store_channel, "BLINKIT")
        self.assertEqual(res.store_region, "NORTH")
        self.assertEqual(res.store_entity, "BCPL")
        self.assertEqual(res.delivery_manager_id, self.dm.id)
        self.assertEqual(res.delivery_manager_name, "Delivery Manager North")
        self.assertTrue(res.compliance_number.startswith("SMC-202610-"))

    def test_create_monthly_compliance_uniqueness(self):
        payload = StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10)
        create_monthly_compliance_endpoint(payload=payload, db=self.db, current_user=self.superadmin)

        with self.assertRaises(HTTPException) as ctx:
            create_monthly_compliance_endpoint(payload=payload, db=self.db, current_user=self.superadmin)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("already exists", str(ctx.exception.detail))

    def test_create_monthly_compliance_different_month_same_store(self):
        p1 = StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10)
        p2 = StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=11)
        r1 = create_monthly_compliance_endpoint(payload=p1, db=self.db, current_user=self.superadmin)
        r2 = create_monthly_compliance_endpoint(payload=p2, db=self.db, current_user=self.superadmin)
        self.assertEqual(r1.month, 10)
        self.assertEqual(r2.month, 11)

    def test_create_monthly_compliance_invalid_month(self):
        with self.assertRaises(Exception):
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=13)

    def test_create_monthly_compliance_invalid_year(self):
        with self.assertRaises(Exception):
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=1999, month=5)

    def test_create_monthly_compliance_nonexistent_store(self):
        payload = StoreMonthlyComplianceCreate(store_id="NON-EXISTENT", year=2026, month=10)
        with self.assertRaises(HTTPException) as ctx:
            create_monthly_compliance_endpoint(payload=payload, db=self.db, current_user=self.superadmin)
        self.assertEqual(ctx.exception.status_code, 404)

    def test_create_monthly_compliance_employee_assigned_store(self):
        payload = StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10)
        res = create_monthly_compliance_endpoint(payload=payload, db=self.db, current_user=self.employee)
        self.assertEqual(res.store_id, "STR-GGN-01")

    def test_create_monthly_compliance_employee_unassigned_store_forbidden(self):
        payload = StoreMonthlyComplianceCreate(store_id="STR-BLR-02", year=2026, month=10)
        with self.assertRaises(HTTPException) as ctx:
            create_monthly_compliance_endpoint(payload=payload, db=self.db, current_user=self.employee)
        self.assertEqual(ctx.exception.status_code, 403)

    def test_create_monthly_compliance_unauthorized_role_forbidden(self):
        payload = StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10)
        with self.assertRaises(HTTPException) as ctx:
            create_monthly_compliance_endpoint(payload=payload, db=self.db, current_user=self.partner)
        self.assertEqual(ctx.exception.status_code, 403)

    # =========================================================================
    # 3. CARD UPLOAD & REPLACEMENT AUDIT (EXACT ATTACHMENT COUNT)
    # =========================================================================

    async def test_upload_monthly_card_success(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(c.card_status, "PENDING_CARD")
        img_bytes = _generate_test_image("green")
        upload = UploadFile(filename="card_oct.jpg", file=io.BytesIO(img_bytes), headers={"content-type": "image/jpeg"})

        res = await upload_monthly_card_endpoint(
            compliance_id=c.id,
            file=upload,
            remarks="Verified stamp attached",
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(res.card_status, "CARD_UPLOADED")
        self.assertIsNotNone(res.card_file_url)
        self.assertEqual(res.card_remarks, "Verified stamp attached")

        # Verify exactly 1 active CARD attachment created
        attachments = self.db.query(StoreComplianceAttachment).filter(
            StoreComplianceAttachment.monthly_compliance_id == c.id,
            StoreComplianceAttachment.attachment_type == StoreComplianceAttachmentType.CARD.value,
        ).all()
        self.assertEqual(len(attachments), 1)
        self.assertIsNone(attachments[0].voided_at)

    async def test_three_card_replacements_produce_exactly_3_card_attachment_records(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )

        # 1. First upload
        img1 = _generate_test_image("red")
        await upload_monthly_card_endpoint(
            compliance_id=c.id,
            file=UploadFile(filename="card_v1.jpg", file=io.BytesIO(img1), headers={"content-type": "image/jpeg"}),
            remarks="Card v1",
            db=self.db,
            current_user=self.superadmin,
        )

        # 2. Second upload (replacement 1)
        img2 = _generate_test_image("blue")
        await upload_monthly_card_endpoint(
            compliance_id=c.id,
            file=UploadFile(filename="card_v2.jpg", file=io.BytesIO(img2), headers={"content-type": "image/jpeg"}),
            remarks="Card v2",
            db=self.db,
            current_user=self.superadmin,
        )

        # 3. Third upload (replacement 2)
        img3 = _generate_test_image("green")
        res3 = await upload_monthly_card_endpoint(
            compliance_id=c.id,
            file=UploadFile(filename="card_v3.jpg", file=io.BytesIO(img3), headers={"content-type": "image/jpeg"}),
            remarks="Card v3 final",
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(res3.card_file_name, "card_v3.jpg")

        # Verify exactly 3 CARD attachments total: 2 voided, 1 active
        attachments = self.db.query(StoreComplianceAttachment).filter(
            StoreComplianceAttachment.monthly_compliance_id == c.id,
            StoreComplianceAttachment.attachment_type == StoreComplianceAttachmentType.CARD.value,
        ).order_by(StoreComplianceAttachment.created_at.asc()).all()

        self.assertEqual(len(attachments), 3, "Expected exactly 3 CARD attachments after 3 uploads")

        voided = [a for a in attachments if a.voided_at is not None]
        active = [a for a in attachments if a.voided_at is None]

        self.assertEqual(len(voided), 2, "Expected exactly 2 voided historical card attachments")
        self.assertEqual(len(active), 1, "Expected exactly 1 active current card attachment")
        self.assertEqual(active[0].file_name, "card_v3.jpg")
        for v in voided:
            self.assertIn("Replaced by new monthly water card upload", v.void_reason)

    async def test_upload_card_authorized_by_existing_store_compliance_upload_permission(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        img = _generate_test_image("purple")
        res = await upload_monthly_card_endpoint(
            compliance_id=c.id,
            file=UploadFile(filename="card_emp.jpg", file=io.BytesIO(img), headers={"content-type": "image/jpeg"}),
            db=self.db,
            current_user=self.employee,
        )
        self.assertEqual(res.card_status, "CARD_UPLOADED")

    async def test_upload_card_invalid_image_corrupted(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        fake_bytes = b"not an image file content"
        upload = UploadFile(filename="corrupted.png", file=io.BytesIO(fake_bytes), headers={"content-type": "image/png"})

        with self.assertRaises(HTTPException) as ctx:
            await upload_monthly_card_endpoint(
                compliance_id=c.id,
                file=upload,
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("corrupted", str(ctx.exception.detail).lower())

    async def test_upload_card_unsupported_extension(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        upload = UploadFile(filename="malicious.exe", file=io.BytesIO(b"executable data"), headers={"content-type": "application/x-msdownload"})

        with self.assertRaises(HTTPException) as ctx:
            await upload_monthly_card_endpoint(
                compliance_id=c.id,
                file=upload,
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("unsupported", str(ctx.exception.detail).lower())

    async def test_upload_card_exceeds_max_size_limit(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        big_bytes = b"x" * (10 * 1024 * 1024 + 1)
        upload = UploadFile(filename="huge.pdf", file=io.BytesIO(big_bytes), headers={"content-type": "application/pdf"})

        with self.assertRaises(HTTPException) as ctx:
            await upload_monthly_card_endpoint(
                compliance_id=c.id,
                file=upload,
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("exceeds", str(ctx.exception.detail).lower())

    async def test_upload_card_employee_scoped(self):
        c1 = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        c2 = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-BLR-02", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        img = _generate_test_image("yellow")

        # Allowed for assigned Store 1
        res1 = await upload_monthly_card_endpoint(
            compliance_id=c1.id,
            file=UploadFile(filename="card1.jpg", file=io.BytesIO(img), headers={"content-type": "image/jpeg"}),
            db=self.db,
            current_user=self.employee,
        )
        self.assertEqual(res1.card_status, "CARD_UPLOADED")

        # Forbidden for unassigned Store 2
        with self.assertRaises(HTTPException) as ctx:
            await upload_monthly_card_endpoint(
                compliance_id=c2.id,
                file=UploadFile(filename="card2.jpg", file=io.BytesIO(img), headers={"content-type": "image/jpeg"}),
                db=self.db,
                current_user=self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)

    # =========================================================================
    # 4. COMMERCIAL INVOICE TESTS (MANDATORY FILE, REPLACEMENT, IMMUTABILITY)
    # =========================================================================

    async def test_invoice_creation_requires_file(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        empty_file = UploadFile(filename="empty.pdf", file=io.BytesIO(b""), headers={"content-type": "application/pdf"})
        with self.assertRaises(HTTPException) as ctx:
            await create_compliance_invoice_endpoint(
                compliance_id=c.id,
                file=empty_file,
                invoice_number="INV-NO-FILE",
                invoice_date=date(2026, 10, 1),
                invoice_amount=Decimal("1000.00"),
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("empty", str(ctx.exception.detail).lower())

    async def test_create_invoice_success_superadmin(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="invoice.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        res = await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-2026-10-001",
            invoice_date=date(2026, 10, 5),
            invoice_amount=Decimal("4500.00"),
            remarks="Monthly delivery invoice",
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(res.invoice_number, "INV-2026-10-001")
        self.assertEqual(res.invoice_amount, Decimal("4500.00"))
        self.assertEqual(res.total_paid, Decimal("0.00"))
        self.assertEqual(res.outstanding_amount, Decimal("4500.00"))
        self.assertEqual(res.payment_status, "UNPAID")
        self.assertIsNotNone(res.file_url)

    async def test_create_invoice_duplicate_prevented(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        f1 = UploadFile(filename="inv1.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=f1,
            invoice_number="INV-001",
            invoice_date=date(2026, 10, 5),
            invoice_amount=Decimal("1000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        f2 = UploadFile(filename="inv2.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        with self.assertRaises(HTTPException) as ctx:
            await create_compliance_invoice_endpoint(
                compliance_id=c.id,
                file=f2,
                invoice_number="INV-002",
                invoice_date=date(2026, 10, 5),
                invoice_amount=Decimal("2000.00"),
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("already attached", str(ctx.exception.detail))

    async def test_create_invoice_non_superadmin_forbidden(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        f = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        with self.assertRaises(HTTPException) as ctx:
            await create_compliance_invoice_endpoint(
                compliance_id=c.id,
                file=f,
                invoice_number="INV-001",
                invoice_date=date(2026, 10, 5),
                invoice_amount=Decimal("1000.00"),
                db=self.db,
                current_user=self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)

    async def test_invoice_replacement_before_payment_preserves_previous_file_evidence(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        f1 = UploadFile(filename="invoice_v1.jpg", file=io.BytesIO(_generate_test_image("red")), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=f1,
            invoice_number="INV-REPLACE",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("3000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        f2 = UploadFile(filename="invoice_v2.jpg", file=io.BytesIO(_generate_test_image("blue")), headers={"content-type": "image/jpeg"})
        updated = await update_compliance_invoice_endpoint(
            compliance_id=c.id,
            invoice_number="INV-REPLACE-2",
            invoice_amount=Decimal("3200.00"),
            file=f2,
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(updated.invoice_number, "INV-REPLACE-2")
        self.assertEqual(updated.invoice_amount, Decimal("3200.00"))
        self.assertEqual(updated.file_name, "invoice_v2.jpg")

        # Verify old invoice document was preserved as soft-voided attachment
        attachments = self.db.query(StoreComplianceAttachment).filter(
            StoreComplianceAttachment.invoice_id == updated.id,
            StoreComplianceAttachment.attachment_type == StoreComplianceAttachmentType.INVOICE.value,
        ).order_by(StoreComplianceAttachment.created_at.asc()).all()

        self.assertEqual(len(attachments), 2)
        self.assertIsNotNone(attachments[0].voided_at)
        self.assertIn("Replaced by new invoice document upload", attachments[0].void_reason)
        self.assertIsNone(attachments[1].voided_at)
        self.assertEqual(attachments[1].file_name, "invoice_v2.jpg")

    async def test_invoice_financial_identity_immutable_after_payment(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        f = UploadFile(filename="inv_lock.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=f,
            invoice_number="INV-LOCK",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("5000.00"),
            db=self.db,
            current_user=self.superadmin,
        )
        record_compliance_payment_endpoint(
            compliance_id=c.id,
            payload=StoreCompliancePaymentCreate(
                amount=Decimal("2000.00"),
                payment_date=datetime.now(timezone.utc),
                payment_mode="UPI",
            ),
            db=self.db,
            current_user=self.superadmin,
        )

        # 1. Attempting to modify amount is blocked
        with self.assertRaises(HTTPException) as ctx:
            await update_compliance_invoice_endpoint(
                compliance_id=c.id,
                invoice_amount=Decimal("6000.00"),
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("immutable", str(ctx.exception.detail).lower())

        # 2. Attempting to replace file after payment is blocked
        f_attempt = UploadFile(filename="inv_attempt.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        with self.assertRaises(HTTPException) as ctx:
            await update_compliance_invoice_endpoint(
                compliance_id=c.id,
                file=f_attempt,
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("immutable", str(ctx.exception.detail).lower())

        # 3. Updating remarks is permitted
        res = await update_compliance_invoice_endpoint(
            compliance_id=c.id,
            remarks="Legitimate note update",
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(res.remarks, "Legitimate note update")

    # =========================================================================
    # 5. PAYMENT LEDGER & CONCURRENT IDENTIFIERS
    # =========================================================================

    async def test_safe_concurrent_payment_number_generation(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-PAY-ID",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("10000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        p1 = record_compliance_payment_endpoint(
            compliance_id=c.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("1000.00"), payment_date=datetime.now(timezone.utc), transaction_reference="REF1"),
            db=self.db,
            current_user=self.superadmin,
        )
        p2 = record_compliance_payment_endpoint(
            compliance_id=c.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("1000.00"), payment_date=datetime.now(timezone.utc), transaction_reference="REF2"),
            db=self.db,
            current_user=self.superadmin,
        )

        self.assertNotEqual(p1.payment_number, p2.payment_number)
        self.assertTrue(p1.payment_number.startswith("CPAY-"))
        self.assertTrue(p2.payment_number.startswith("CPAY-"))
        self.assertEqual(p1.payment_number, f"CPAY-{p1.id:07d}")
        self.assertEqual(p2.payment_number, f"CPAY-{p2.id:07d}")

    async def test_record_payment_partial_and_full(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-SPLIT",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("5000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        # Partial
        record_compliance_payment_endpoint(
            compliance_id=c.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("2000.00"), payment_date=datetime.now(timezone.utc)),
            db=self.db,
            current_user=self.superadmin,
        )
        inv = self.db.query(StoreComplianceInvoice).filter(StoreComplianceInvoice.monthly_compliance_id == c.id).first()
        self.assertEqual(inv.total_paid, Decimal("2000.00"))
        self.assertEqual(inv.outstanding_amount, Decimal("3000.00"))
        self.assertEqual(inv.payment_status, "PARTIALLY_PAID")

        # Full
        record_compliance_payment_endpoint(
            compliance_id=c.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("3000.00"), payment_date=datetime.now(timezone.utc)),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_paid = self.db.query(StoreComplianceInvoice).filter(StoreComplianceInvoice.monthly_compliance_id == c.id).first()
        self.assertEqual(inv_paid.total_paid, Decimal("5000.00"))
        self.assertEqual(inv_paid.outstanding_amount, Decimal("0.00"))
        self.assertEqual(inv_paid.payment_status, "PAID")

    async def test_record_payment_overpayment_rejected(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-OVER",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("1000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        with self.assertRaises(HTTPException) as ctx:
            record_compliance_payment_endpoint(
                compliance_id=c.id,
                payload=StoreCompliancePaymentCreate(amount=Decimal("1500.00"), payment_date=datetime.now(timezone.utc)),
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("exceeds", str(ctx.exception.detail).lower())

    async def test_record_payment_non_superadmin_forbidden(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-FORBID",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("1000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        with self.assertRaises(HTTPException) as ctx:
            record_compliance_payment_endpoint(
                compliance_id=c.id,
                payload=StoreCompliancePaymentCreate(amount=Decimal("500.00"), payment_date=datetime.now(timezone.utc)),
                db=self.db,
                current_user=self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)

    # =========================================================================
    # 6. ATTACHMENT PERMISSIONS & SOFT VOIDING
    # =========================================================================

    async def test_employee_cannot_upload_invoice_attachment(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        inv = await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-ATT",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("1000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        doc = UploadFile(filename="doc.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        with self.assertRaises(HTTPException) as ctx:
            await upload_compliance_attachment_endpoint(
                compliance_id=c.id,
                file=doc,
                attachment_type="INVOICE",
                invoice_id=inv.id,
                db=self.db,
                current_user=self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)

    async def test_employee_cannot_upload_payment_proof(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        inv_file = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        inv = await create_compliance_invoice_endpoint(
            compliance_id=c.id,
            file=inv_file,
            invoice_number="INV-PROOF",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("1000.00"),
            db=self.db,
            current_user=self.superadmin,
        )
        pay = record_compliance_payment_endpoint(
            compliance_id=c.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("500.00"), payment_date=datetime.now(timezone.utc)),
            db=self.db,
            current_user=self.superadmin,
        )

        proof = UploadFile(filename="utr.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        with self.assertRaises(HTTPException) as ctx:
            await upload_compliance_attachment_endpoint(
                compliance_id=c.id,
                file=proof,
                attachment_type="PAYMENT_PROOF",
                invoice_id=inv.id,
                payment_id=pay.id,
                db=self.db,
                current_user=self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)

    async def test_employee_can_upload_other_supporting_attachment(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        doc = UploadFile(filename="challan.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        att = await upload_compliance_attachment_endpoint(
            compliance_id=c.id,
            file=doc,
            attachment_type="OTHER",
            db=self.db,
            current_user=self.employee,
        )
        self.assertEqual(att.attachment_type, "OTHER")

    async def test_employee_cannot_void_attachment(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        doc = UploadFile(filename="doc.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        att = await upload_compliance_attachment_endpoint(
            compliance_id=c.id,
            file=doc,
            attachment_type="OTHER",
            db=self.db,
            current_user=self.superadmin,
        )

        with self.assertRaises(HTTPException) as ctx:
            void_compliance_attachment_endpoint(
                compliance_id=c.id,
                attachment_id=att.id,
                payload=StoreComplianceAttachmentVoidRequest(void_reason="Mistake"),
                db=self.db,
                current_user=self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)

    async def test_void_attachment_superadmin_only(self):
        c = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        doc = UploadFile(filename="doc.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        att = await upload_compliance_attachment_endpoint(
            compliance_id=c.id,
            file=doc,
            attachment_type="OTHER",
            db=self.db,
            current_user=self.superadmin,
        )

        voided = void_compliance_attachment_endpoint(
            compliance_id=c.id,
            attachment_id=att.id,
            payload=StoreComplianceAttachmentVoidRequest(void_reason="Uploaded incorrect scan"),
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertIsNotNone(voided.voided_at)
        self.assertEqual(voided.void_reason, "Uploaded incorrect scan")

    # =========================================================================
    # 7. SUMMARY & EXCEL EXPORT
    # =========================================================================

    async def test_summary_kpis_and_aggregations(self):
        c1 = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        await upload_monthly_card_endpoint(
            compliance_id=c1.id,
            file=UploadFile(filename="card1.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"}),
            db=self.db,
            current_user=self.superadmin,
        )
        f1 = UploadFile(filename="inv1.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c1.id,
            file=f1,
            invoice_number="INV-SUM-1",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("3000.00"),
            db=self.db,
            current_user=self.superadmin,
        )
        record_compliance_payment_endpoint(
            compliance_id=c1.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("3000.00"), payment_date=datetime.now(timezone.utc)),
            db=self.db,
            current_user=self.superadmin,
        )

        c2 = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-BLR-02", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        f2 = UploadFile(filename="inv2.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c2.id,
            file=f2,
            invoice_number="INV-SUM-2",
            invoice_date=date(2026, 10, 2),
            invoice_amount=Decimal("5000.00"),
            db=self.db,
            current_user=self.superadmin,
        )
        record_compliance_payment_endpoint(
            compliance_id=c2.id,
            payload=StoreCompliancePaymentCreate(amount=Decimal("2000.00"), payment_date=datetime.now(timezone.utc)),
            db=self.db,
            current_user=self.superadmin,
        )

        summary = get_monthly_compliance_summary_endpoint(year=2026, month=10, db=self.db, current_user=self.superadmin)
        self.assertEqual(summary.total_monthly_records, 2)
        self.assertEqual(summary.cards_uploaded, 1)
        self.assertEqual(summary.pending_cards, 1)
        self.assertEqual(summary.with_invoice, 2)
        self.assertEqual(summary.without_invoice, 0)
        self.assertEqual(summary.paid_count, 1)
        self.assertEqual(summary.partially_paid_count, 1)
        self.assertEqual(summary.unpaid_count, 0)
        self.assertEqual(summary.total_invoice_amount, Decimal("8000.00"))
        self.assertEqual(summary.total_paid_amount, Decimal("5000.00"))
        self.assertEqual(summary.total_outstanding_amount, Decimal("3000.00"))

    async def test_excel_export_and_filtering(self):
        c1 = create_monthly_compliance_endpoint(
            StoreMonthlyComplianceCreate(store_id="STR-GGN-01", year=2026, month=10),
            db=self.db,
            current_user=self.superadmin,
        )
        f = UploadFile(filename="inv.jpg", file=io.BytesIO(_generate_test_image()), headers={"content-type": "image/jpeg"})
        await create_compliance_invoice_endpoint(
            compliance_id=c1.id,
            file=f,
            invoice_number="INV-EXP-1",
            invoice_date=date(2026, 10, 1),
            invoice_amount=Decimal("4000.00"),
            db=self.db,
            current_user=self.superadmin,
        )

        res = export_monthly_compliance_excel_endpoint(year=2026, month=10, db=self.db, current_user=self.superadmin)
        self.assertEqual(res.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("Content-Disposition", res.headers)

        wb = openpyxl.load_workbook(io.BytesIO(res.body))
        ws = wb.active
        self.assertEqual(ws.title, "Store Compliance Monthly")
        self.assertEqual(ws.cell(row=1, column=1).value, "Compliance No")
        self.assertEqual(ws.cell(row=2, column=2).value, "STR-GGN-01")
        self.assertEqual(ws.cell(row=2, column=10).value, "INV-EXP-1")

        list_res = list_monthly_compliances_endpoint(
            search="EXP-1",
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(list_res.total, 1)
        self.assertEqual(list_res.items[0].store_id, "STR-GGN-01")

        detail = get_monthly_compliance_detail_endpoint(
            compliance_id=c1.id,
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertIsNotNone(detail.invoice)
        self.assertEqual(detail.invoice.invoice_number, "INV-EXP-1")


if __name__ == "__main__":
    unittest.main()
