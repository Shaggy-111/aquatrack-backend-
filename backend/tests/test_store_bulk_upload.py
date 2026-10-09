import io
import os
import unittest
from datetime import datetime, timezone
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from fastapi import HTTPException, UploadFile

import backend.models
from backend.database.db import Base
from backend.models.store import Store
from backend.models.user import User, RoleEnum
from backend.models.order import Order, OrderStatusEnum
from backend.models.partner_store import PartnerStore
from backend.schemas.store import StoreCreate, StoreUpdate
from backend.routes.store import (
    bulk_upload_stores,
    download_bulk_upload_template,
    create_store,
    get_stores,
    update_store,
)
from backend.utils.store_validation import (
    ALL_EXCEL_HEADERS,
    generate_bulk_upload_template,
)


def make_excel_bytes(rows, headers=None, sheet_name="Stores"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    h = headers if headers is not None else ALL_EXCEL_HEADERS
    ws.append(h)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class StoreBulkUploadTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        os.environ["STORE_BULK_UPLOAD_PASSWORD"] = "SecretBulkPass123"

        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()

        self.admin = User(
            id=1,
            email="bulk-admin@test.local",
            password="test",
            full_name="Super Admin",
            role=RoleEnum.superadmin,
            status="active",
        )
        self.customer = User(
            id=2,
            email="bulk-customer@test.local",
            password="test",
            full_name="Customer User",
            role=RoleEnum.customer,
            status="active",
        )
        self.db.add_all([self.admin, self.customer])
        self.db.commit()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()
        if "STORE_BULK_UPLOAD_PASSWORD" in os.environ:
            del os.environ["STORE_BULK_UPLOAD_PASSWORD"]

    def _make_upload_file(self, content_bytes, filename="stores.xlsx"):
        file_obj = io.BytesIO(content_bytes)
        return UploadFile(filename=filename, file=file_obj)

    # 1. Valid Excel creates all Stores.
    async def test_valid_excel_creates_all_stores(self):
        rows = [
            ["STR-B01", "Bulk Store 1", "NORTH", "BCPL", "Haryana", "Gurgaon", "Addr 1", "122001", 28.45, 77.04, "BLINKIT", "P-1", "POC One", "9876543210", "TRUE"],
            ["STR-B02", "Bulk Store 2", "SOUTH", "BISTRO", "Tamil Nadu", "Chennai", "Addr 2", "600001", 13.08, 80.27, "ZEPTO", "P-2", "POC Two", "9876543211", "TRUE"],
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        self.assertTrue(res["success"])
        self.assertEqual(res["created_count"], 2)
        self.assertEqual(res["total_rows"], 2)
        self.assertEqual(self.db.query(Store).count(), 2)

    # 2. Wrong PIN/password rejected.
    async def test_wrong_password_rejected(self):
        excel_bytes = make_excel_bytes([["STR-01", "S1", "NORTH", "BCPL", "Delhi", "Delhi", "A", "110001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]])
        upload_file = self._make_upload_file(excel_bytes)

        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password="WrongPassword",
                db=self.db,
                current_user=self.admin
            )
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("Invalid bulk upload password", ctx.exception.detail)

    # 3. Missing PIN/password rejected.
    async def test_missing_password_rejected(self):
        excel_bytes = make_excel_bytes([])
        upload_file = self._make_upload_file(excel_bytes)

        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password=None,
                pin=None,
                db=self.db,
                current_user=self.admin
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("password/PIN is required", ctx.exception.detail)

    # 4. Non-admin rejected.
    async def test_non_admin_rejected(self):
        excel_bytes = make_excel_bytes([])
        upload_file = self._make_upload_file(excel_bytes)

        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password="SecretBulkPass123",
                db=self.db,
                current_user=self.customer
            )
        self.assertEqual(ctx.exception.status_code, 403)

    # 5. Wrong Excel type rejected.
    async def test_wrong_excel_type_rejected(self):
        upload_file = self._make_upload_file(b"csv,data", filename="stores.csv")
        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password="SecretBulkPass123",
                db=self.db,
                current_user=self.admin
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("Only .xlsx Excel files are supported", ctx.exception.detail)

    # 6. Missing required header rejected.
    async def test_missing_required_header_rejected(self):
        headers = ["STORE_NAME", "CHANNEL"]  # Missing STORE_ID
        excel_bytes = make_excel_bytes([["Name", "BLINKIT"]], headers=headers)
        upload_file = self._make_upload_file(excel_bytes)

        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password="SecretBulkPass123",
                db=self.db,
                current_user=self.admin
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("Missing required header", ctx.exception.detail)

    # 7. Duplicate header rejected.
    async def test_duplicate_header_rejected(self):
        headers = ["STORE_ID", "STORE_NAME", "CHANNEL", "STORE_ID"]
        excel_bytes = make_excel_bytes([["S1", "Name", "BLINKIT", "S1"]], headers=headers)
        upload_file = self._make_upload_file(excel_bytes)

        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password="SecretBulkPass123",
                db=self.db,
                current_user=self.admin
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("Duplicate header found", ctx.exception.detail)

    # 8. Wrong-case master value rejected.
    async def test_wrong_case_master_value_rejected(self):
        rows = [
            ["STR-C01", "Store Case 1", "North", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        # Should return JSONResponse with status 400 and structured errors
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        self.assertEqual(body["invalid_rows"], 1)
        self.assertIn("Invalid REGION", body["errors"][0]["message"])

    # 9. Duplicate Store ID within workbook rejected.
    async def test_duplicate_store_id_in_workbook_rejected(self):
        rows = [
            ["STR-DUP", "Store One", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"],
            ["STR-DUP", "Store Two", "NORTH", "BCPL", "Haryana", "Gurgaon", "B", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543211", "TRUE"],
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        self.assertTrue(any("Duplicate STORE_ID" in e["message"] for e in body["errors"]))

    # 10. Existing Store ID rejected.
    async def test_existing_store_id_rejected(self):
        # Pre-seed store
        self.db.add(Store(id="EXIST-01", store_name="Existing Store", channel="GENERAL"))
        self.db.commit()

        rows = [
            ["EXIST-01", "New Store Same ID", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        self.assertTrue(any("already exists in database" in e["message"] for e in body["errors"]))

    # 11. Invalid numeric value rejected.
    async def test_invalid_numeric_value_rejected(self):
        rows = [
            ["STR-NUM", "Store Num", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", "invalid_lat", 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        self.assertTrue(any("Invalid LATITUDE" in e["message"] for e in body["errors"]))

    # 12. Invalid required field rejected.
    async def test_invalid_required_field_rejected(self):
        rows = [
            ["", "Store No ID", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        self.assertTrue(any("STORE_ID is required" in e["message"] for e in body["errors"]))

    # 13. One bad row causes zero inserts.
    async def test_one_bad_row_causes_zero_inserts(self):
        rows = [
            ["GOOD-01", "Good Store", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"],
            ["BAD-02", "Bad Store", "INVALID_REG", "BCPL", "Haryana", "Gurgaon", "B", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543211", "TRUE"],
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        # Ensure zero inserts in DB
        self.assertEqual(self.db.query(Store).filter(Store.id == "GOOD-01").count(), 0)

    # 14. Valid canonical values succeed.
    async def test_valid_canonical_values_succeed(self):
        rows = [
            ["CAN-01", "Canonical Store", "CENTRAL", "ZHPL", "Madhya Pradesh", "Bhopal", "Main Rd", "462001", 23.25, 77.41, "IBM", "P-IBM", "POC Name", "9876543210", "TRUE"]
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(
            file=upload_file,
            password="SecretBulkPass123",
            db=self.db,
            current_user=self.admin
        )
        self.assertTrue(res["success"])
        self.assertEqual(res["created_count"], 1)

    # 15. Template endpoint works.
    def test_template_endpoint_works(self):
        res = download_bulk_upload_template(current_user=self.admin)
        self.assertEqual(res.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("store_bulk_upload_template.xlsx", res.headers["Content-Disposition"])
        # Parse returned workbook
        wb = openpyxl.load_workbook(io.BytesIO(res.body))
        self.assertIn("Stores", wb.sheetnames)
        self.assertIn("Allowed Values", wb.sheetnames)
        ws_stores = wb["Stores"]
        headers = [c.value for c in ws_stores[1]]
        self.assertEqual(headers, ALL_EXCEL_HEADERS)

    # 16. Existing Add Store still works.
    def test_existing_add_store_still_works(self):
        store_in = StoreCreate(
            id="ADD-01",
            store_name="Direct Added Store",
            channel="GENERAL",
            city="Kolkata",
            state="West Bengal",
            region="EAST",
        )
        res = create_store(store_in, self.db, self.admin)
        self.assertEqual(res.id, "ADD-01")
        self.assertEqual(res.store_name, "Direct Added Store")

    # 17. Existing Store list still works.
    def test_existing_store_list_still_works(self):
        self.db.add(Store(id="LIST-01", store_name="List Store", channel="GENERAL", is_active=True))
        self.db.commit()
        stores = get_stores(self.db, self.admin)
        self.assertTrue(any(s.id == "LIST-01" for s in stores))

    # 18. Existing Store update still works.
    def test_existing_store_update_still_works(self):
        self.db.add(Store(id="UPD-01", store_name="Before Update", channel="GENERAL", is_active=True))
        self.db.commit()
        updated = update_store("UPD-01", StoreUpdate(store_name="After Update"), self.db, self.admin)
        self.assertEqual(updated.store_name, "After Update")

    # 19. Historical records remain intact.
    async def test_historical_records_remain_intact(self):
        # Seed an existing store and order
        partner = User(id=10, email="p10@test.local", password="t", full_name="P10", role=RoleEnum.partner, status="active")
        store = Store(id="HIST-01", store_name="Hist Store", channel="BLINKIT")
        self.db.add_all([partner, store])
        self.db.flush()
        order = Order(
            id=500,
            store_id=store.id,
            partner_id=partner.id,
            channel="BLINKIT",
            order_details="5 jars",
            total_amount=250.0,
            status=OrderStatusEnum.delivered_confirmed,
            poc_name="Old Partner",
        )
        self.db.add(order)
        self.db.commit()

        # Run bulk upload for completely new store
        rows = [
            ["NEW-01", "New Store", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)
        res = await bulk_upload_stores(upload_file, password="SecretBulkPass123", db=self.db, current_user=self.admin)
        self.assertTrue(res["success"])

        # Check historical order is intact
        hist_order = self.db.query(Order).filter(Order.id == 500).first()
        self.assertEqual(hist_order.poc_name, "Old Partner")
        self.assertEqual(hist_order.store_id, "HIST-01")
        self.assertEqual(hist_order.total_amount, 250.0)

    # 20. Employee cannot bulk upload (rejected with 403)
    async def test_employee_cannot_bulk_upload(self):
        employee = User(
            id=20,
            email="bulk-emp@test.local",
            password="test",
            full_name="Bulk Employee",
            role=RoleEnum.employee,
            status="active",
        )
        self.db.add(employee)
        self.db.commit()

        excel_bytes = make_excel_bytes([["STR-E01", "S1", "NORTH", "BCPL", "Delhi", "Delhi", "A", "110001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]])
        upload_file = self._make_upload_file(excel_bytes)

        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password="SecretBulkPass123",
                db=self.db,
                current_user=employee
            )
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("Only an admin or super admin", ctx.exception.detail)

    # 21. Employee cannot download template (rejected with 403)
    def test_employee_cannot_download_template(self):
        employee = User(
            id=21,
            email="tmpl-emp@test.local",
            password="test",
            full_name="Template Employee",
            role=RoleEnum.employee,
            status="active",
        )
        self.db.add(employee)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            download_bulk_upload_template(current_user=employee)
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("Only an admin or super admin", ctx.exception.detail)

    # 22. Admin role can bulk upload and download template
    async def test_admin_role_can_bulk_upload_and_download_template(self):
        admin_user = User(
            id=22,
            email="reg-admin@test.local",
            password="test",
            full_name="Regular Admin User",
            role=RoleEnum.admin,
            status="active",
        )
        self.db.add(admin_user)
        self.db.commit()

        # Download template
        res_tmpl = download_bulk_upload_template(current_user=admin_user)
        self.assertEqual(res_tmpl.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        # Bulk upload
        rows = [
            ["ADM-01", "Admin Store", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)
        res = await bulk_upload_stores(upload_file, password="SecretBulkPass123", db=self.db, current_user=admin_user)
        self.assertTrue(res["success"])
        self.assertEqual(res["created_count"], 1)

    # 23. Missing server password fails closed (500)
    async def test_missing_server_bulk_upload_password_fails_closed(self):
        if "STORE_BULK_UPLOAD_PASSWORD" in os.environ:
            del os.environ["STORE_BULK_UPLOAD_PASSWORD"]

        excel_bytes = make_excel_bytes([["STR-01", "S1", "NORTH", "BCPL", "Delhi", "Delhi", "A", "110001", 28.0, 77.0, "BLINKIT", "P", "POC", "9876543210", "TRUE"]])
        upload_file = self._make_upload_file(excel_bytes)

        with self.assertRaises(HTTPException) as ctx:
            await bulk_upload_stores(
                file=upload_file,
                password="AnyPassword",
                db=self.db,
                current_user=self.admin
            )
        self.assertEqual(ctx.exception.status_code, 500)
        self.assertIn("Bulk upload is disabled", ctx.exception.detail)

    # 24. Multi-store same POC phone reuses the same partner User
    async def test_multi_store_same_poc_phone_reuses_partner(self):
        rows = [
            ["STR-S1", "Store One", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P1", "Shared POC", "9555444333", "TRUE"],
            ["STR-S2", "Store Two", "SOUTH", "ZEPTO", "Karnataka", "Bangalore", "B", "560001", 12.9, 77.5, "ZEPTO", "P2", "Shared POC", "9555444333", "TRUE"],
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(upload_file, password="SecretBulkPass123", db=self.db, current_user=self.admin)
        self.assertTrue(res["success"])
        self.assertEqual(res["created_count"], 2)

        # Check only 1 partner user was created with this phone
        partners = self.db.query(User).filter(User.mobile_number == "9555444333").all()
        self.assertEqual(len(partners), 1)
        partner_id = partners[0].id

        # Check both stores are mapped to this partner in PartnerStore
        mappings = self.db.query(PartnerStore).filter(PartnerStore.user_id == partner_id).all()
        self.assertEqual(len(mappings), 2)
        mapped_store_ids = {m.store_id for m in mappings}
        self.assertEqual(mapped_store_ids, {"STR-S1", "STR-S2"})

    # 25. Conflicting POC name for same phone in workbook fails validation
    async def test_conflicting_poc_name_for_same_phone_in_workbook_fails(self):
        rows = [
            ["STR-C1", "Store One", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P1", "Alice Smith", "9777888999", "TRUE"],
            ["STR-C2", "Store Two", "NORTH", "BCPL", "Haryana", "Gurgaon", "B", "122001", 28.0, 77.0, "BLINKIT", "P2", "Bob Jones", "9777888999", "TRUE"],
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(upload_file, password="SecretBulkPass123", db=self.db, current_user=self.admin)
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        self.assertTrue(any("conflicting POC_NAMEs within uploaded file" in e["message"] for e in body["errors"]))

    # 26. Non-partner phone collision in bulk upload fails validation
    async def test_non_partner_phone_collision_in_bulk_fails(self):
        # Create an existing delivery manager user with this phone
        dm = User(
            id=30,
            email="dm_phone@test.local",
            password="test",
            full_name="Delivery Manager Phone",
            mobile_number="9333222111",
            role=RoleEnum.delivery_manager,
            status="active",
        )
        self.db.add(dm)
        self.db.commit()

        rows = [
            ["STR-DM1", "Store DM", "NORTH", "BCPL", "Haryana", "Gurgaon", "A", "122001", 28.0, 77.0, "BLINKIT", "P1", "DM Collide", "9333222111", "TRUE"],
        ]
        excel_bytes = make_excel_bytes(rows)
        upload_file = self._make_upload_file(excel_bytes)

        res = await bulk_upload_stores(upload_file, password="SecretBulkPass123", db=self.db, current_user=self.admin)
        import json
        body = json.loads(res.body.decode("utf-8"))
        self.assertFalse(body["success"])
        self.assertTrue(any("already belongs to an existing non-partner user" in e["message"] for e in body["errors"]))
