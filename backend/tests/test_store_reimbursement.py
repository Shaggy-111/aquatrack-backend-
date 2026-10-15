import io
import unittest
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import openpyxl
from fastapi import HTTPException, UploadFile
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import backend.models
from backend.database.db import Base
from backend.models.employee_permission import EmployeePermission
from backend.models.employee_store_mapping import EmployeeStoreMapping
from backend.models.partner_store import PartnerStore
from backend.models.permission import Permission
from backend.models.store import Store
from backend.models.store_reimbursement import (
    StoreReimbursement,
    StoreReimbursementAttachment,
    StoreReimbursementPayment,
    StoreReimbursementStatus,
)
from backend.models.user import RoleEnum, User
from backend.routes.store_reimbursement import (
    approve_reimbursement_endpoint,
    cancel_reimbursement_endpoint,
    create_reimbursement_endpoint,
    export_reimbursements_endpoint,
    get_reimbursement_detail_endpoint,
    get_reimbursement_summary_endpoint,
    list_reimbursements_endpoint,
    record_payment_endpoint,
    reject_reimbursement_endpoint,
    update_reimbursement_endpoint,
    upload_attachment_endpoint,
    void_attachment_endpoint,
)
from backend.schemas.store_reimbursement import (
    StoreReimbursementAttachmentVoidRequest,
    StoreReimbursementCancel,
    StoreReimbursementCreate,
    StoreReimbursementPaymentCreate,
    StoreReimbursementReject,
    StoreReimbursementUpdate,
)
from backend.utils.permission_seeder import build_permission_rows


class StoreReimbursementModuleTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()

        # 1. SuperAdmin User
        self.superadmin = User(
            id=1,
            email="superadmin@aquatrack.local",
            password="test",
            full_name="Global SuperAdmin",
            role=RoleEnum.superadmin,
            status="active",
        )

        # 2. Employee User
        self.employee = User(
            id=2,
            email="emp@aquatrack.local",
            password="test",
            full_name="Operations Employee",
            role=RoleEnum.employee,
            status="active",
        )

        # 3. Another Employee (Unassigned)
        self.other_employee = User(
            id=3,
            email="other_emp@aquatrack.local",
            password="test",
            full_name="Other Employee",
            role=RoleEnum.employee,
            status="active",
        )

        # 4. Partner User 1
        self.partner_1 = User(
            id=4,
            email="poc1@aquatrack.local",
            password="test",
            full_name="Partner One",
            mobile_number="9876543210",
            role=RoleEnum.partner,
            status="active",
        )

        # 5. Partner User 2
        self.partner_2 = User(
            id=5,
            email="poc2@aquatrack.local",
            password="test",
            full_name="Partner Two",
            mobile_number="9876543211",
            role=RoleEnum.partner,
            status="active",
        )

        # 6. Customer User (Unauthorized role)
        self.customer = User(
            id=6,
            email="customer@aquatrack.local",
            password="test",
            full_name="Customer User",
            role=RoleEnum.customer,
            status="active",
        )

        self.db.add_all([
            self.superadmin,
            self.employee,
            self.other_employee,
            self.partner_1,
            self.partner_2,
            self.customer,
        ])
        self.db.flush()

        # Seed Stores
        self.store_1 = Store(
            id="STR-GGN-01",
            store_name="Gurgaon Sector 29",
            channel="BLINKIT",
            region="NORTH",
            entity="BCPL",
            city="Gurgaon",
            state="Haryana",
            is_active=True,
        )
        self.store_2 = Store(
            id="STR-BLR-02",
            store_name="Bangalore Indiranagar",
            channel="ZEPTO",
            region="SOUTH",
            entity="ZEPTO",
            city="Bangalore",
            state="Karnataka",
            is_active=True,
        )
        self.db.add_all([self.store_1, self.store_2])
        self.db.flush()

        # Link Partner 1 to Store 1 (Single POC)
        self.db.add(PartnerStore(user_id=self.partner_1.id, store_id=self.store_1.id))

        # Assign Store 1 to Employee (Store 2 unassigned)
        self.db.add(
            EmployeeStoreMapping(
                employee_id=self.employee.id,
                store_id=self.store_1.id,
                assigned_by=self.superadmin.id,
            )
        )

        # Seed Permissions
        perm_keys = [
            "store_reimbursement.view",
            "store_reimbursement.create",
            "store_reimbursement.edit",
            "store_reimbursement.approve",
            "store_reimbursement.record_payment",
            "store_reimbursement.upload_proof",
            "store_reimbursement.cancel",
            "store_reimbursement.export",
        ]
        self.perms = {}
        for idx, key in enumerate(perm_keys, start=1):
            p = Permission(
                id=idx,
                module="Store Reimbursement",
                action=key.split(".")[1],
                permission_key=key,
                display_name=key,
                is_active=True,
            )
            self.db.add(p)
            self.perms[key] = p
        self.db.flush()

        # Assign all store_reimbursement permissions to Employee
        for p in self.perms.values():
            self.db.add(
                EmployeePermission(
                    employee_id=self.employee.id,
                    permission_id=p.id,
                    assigned_by=self.superadmin.id,
                )
            )

        self.db.commit()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    def _make_image_bytes(self) -> bytes:
        from PIL import Image
        buf = io.BytesIO()
        img = Image.new("RGB", (50, 50), color="blue")
        img.save(buf, format="PNG")
        return buf.getvalue()

    # =========================================================================
    # 1. CREATION TESTS
    # =========================================================================

    def test_superadmin_can_create_reimbursement_for_any_store(self):
        payload = StoreReimbursementCreate(
            store_id=self.store_2.id,
            purchase_date=date(2026, 10, 10),
            total_amount=Decimal("1500.00"),
            reason="Water delivery truck breakdown; authorized local cans.",
            instructed_by_user_id=self.employee.id,
            notes="Authorized 15 jars @ 100",
        )
        # Link a partner to store_2 first
        self.db.add(PartnerStore(user_id=self.partner_2.id, store_id=self.store_2.id))
        self.db.commit()

        res = create_reimbursement_endpoint(payload, self.db, self.superadmin)
        self.assertEqual(res.store_id, "STR-BLR-02")
        self.assertEqual(res.store_name, "Bangalore Indiranagar")
        self.assertEqual(res.store_channel, "ZEPTO")
        self.assertEqual(res.store_region, "SOUTH")
        self.assertEqual(res.store_entity, "ZEPTO")
        self.assertEqual(res.poc_id, self.partner_2.id)
        self.assertEqual(res.poc_name, "Partner Two")
        self.assertEqual(res.poc_phone, "9876543211")
        self.assertEqual(res.total_amount, Decimal("1500.00"))
        self.assertEqual(res.total_paid, Decimal("0.00"))
        self.assertEqual(res.outstanding_amount, Decimal("1500.00"))
        self.assertEqual(res.status, "PENDING")
        self.assertTrue(res.reimbursement_number.startswith("SR-202610-"))

    def test_employee_can_create_reimbursement_for_assigned_store(self):
        payload = StoreReimbursementCreate(
            store_id=self.store_1.id,
            purchase_date=date(2026, 10, 11),
            total_amount=Decimal("800.00"),
            reason="Plant maintenance; morning shift emergency.",
            instructed_by_user_id=self.employee.id,
        )
        res = create_reimbursement_endpoint(payload, self.db, self.employee)
        self.assertEqual(res.store_id, "STR-GGN-01")
        self.assertEqual(res.total_amount, Decimal("800.00"))
        self.assertEqual(res.poc_name, "Partner One")
        self.assertEqual(res.created_by_id, self.employee.id)

    def test_employee_cannot_create_reimbursement_for_unassigned_store(self):
        payload = StoreReimbursementCreate(
            store_id=self.store_2.id,
            purchase_date=date(2026, 10, 11),
            total_amount=Decimal("500.00"),
            reason="Unassigned test",
        )
        with self.assertRaises(HTTPException) as ctx:
            create_reimbursement_endpoint(payload, self.db, self.employee)
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("not assigned", ctx.exception.detail)

    def test_employee_without_create_permission_rejected(self):
        # Remove create permission from employee
        self.db.query(EmployeePermission).filter(
            EmployeePermission.employee_id == self.employee.id,
            EmployeePermission.permission_id == self.perms["store_reimbursement.create"].id,
        ).delete()
        self.db.commit()

        payload = StoreReimbursementCreate(
            store_id=self.store_1.id,
            purchase_date=date(2026, 10, 11),
            total_amount=Decimal("500.00"),
            reason="Perm check",
        )
        with self.assertRaises(HTTPException) as ctx:
            create_reimbursement_endpoint(payload, self.db, self.employee)
        self.assertEqual(ctx.exception.status_code, 403)

    def test_multiple_poc_ambiguity_requires_explicit_poc_id(self):
        # Link partner_2 also to store_1 so store_1 has 2 partners
        self.db.add(PartnerStore(user_id=self.partner_2.id, store_id=self.store_1.id))
        self.db.commit()

        # Without poc_id -> should fail with clear message
        payload = StoreReimbursementCreate(
            store_id=self.store_1.id,
            purchase_date=date(2026, 10, 11),
            total_amount=Decimal("500.00"),
            reason="Ambiguous POC test",
        )
        with self.assertRaises(HTTPException) as ctx:
            create_reimbursement_endpoint(payload, self.db, self.superadmin)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("Multiple POCs exist", ctx.exception.detail)

        # With explicit poc_id -> succeeds cleanly
        payload.poc_id = self.partner_2.id
        res = create_reimbursement_endpoint(payload, self.db, self.superadmin)
        self.assertEqual(res.poc_id, self.partner_2.id)
        self.assertEqual(res.poc_name, "Partner Two")

    # =========================================================================
    # 2. VIEW / LIST / DETAIL & SCOPING TESTS
    # =========================================================================

    def test_employee_only_sees_assigned_store_claims_in_list(self):
        # Seed 1 claim for store_1 (assigned) and 1 claim for store_2 (unassigned)
        r1 = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("100.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("100.00"),
            status="PENDING",
            reason="R1",
            created_by_id=self.superadmin.id,
        )
        r2 = StoreReimbursement(
            reimbursement_number="SR-02",
            store_id=self.store_2.id,
            store_name="Store 2",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("200.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("200.00"),
            status="PENDING",
            reason="R2",
            created_by_id=self.superadmin.id,
        )
        self.db.add_all([r1, r2])
        self.db.commit()

        # SuperAdmin sees both
        res_admin = list_reimbursements_endpoint(db=self.db, current_user=self.superadmin)
        self.assertEqual(res_admin.total, 2)

        # Employee only sees Store 1 claim
        res_emp = list_reimbursements_endpoint(db=self.db, current_user=self.employee)
        self.assertEqual(res_emp.total, 1)
        self.assertEqual(res_emp.items[0].store_id, "STR-GGN-01")

    def test_employee_cannot_access_unassigned_store_detail(self):
        r2 = StoreReimbursement(
            reimbursement_number="SR-02",
            store_id=self.store_2.id,
            store_name="Store 2",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("200.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("200.00"),
            status="PENDING",
            reason="R2",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r2)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            get_reimbursement_detail_endpoint(r2.id, self.db, self.employee)
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("not assigned", ctx.exception.detail)

    # =========================================================================
    # 3. EDIT TESTS
    # =========================================================================

    def test_pending_claim_is_editable_and_recalculates_outstanding(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Original reason",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        update_payload = StoreReimbursementUpdate(
            total_amount=Decimal("1200.00"),
            reason="Updated revised reason",
        )
        res = update_reimbursement_endpoint(r.id, update_payload, self.db, self.employee)
        self.assertEqual(res.total_amount, Decimal("1200.00"))
        self.assertEqual(res.outstanding_amount, Decimal("1200.00"))
        self.assertEqual(res.reason, "Updated revised reason")

    def test_approved_claim_cannot_be_edited(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="Approved claim",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            update_reimbursement_endpoint(r.id, StoreReimbursementUpdate(reason="Try edit"), self.db, self.superadmin)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("Only PENDING claims can be edited", ctx.exception.detail)

    # =========================================================================
    # 4. APPROVAL & REJECTION TESTS
    # =========================================================================

    def test_superadmin_can_approve_pending_claim(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Approval test",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        res = approve_reimbursement_endpoint(r.id, self.db, self.superadmin)
        self.assertEqual(res.status, "APPROVED")
        self.assertEqual(res.approved_by_id, self.superadmin.id)
        self.assertIsNotNone(res.approved_at)

    def test_employee_cannot_approve_claim(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Approval test",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            approve_reimbursement_endpoint(r.id, self.db, self.employee)
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("Only Super Admin", ctx.exception.detail)

    def test_rejection_requires_valid_reason_and_sets_status(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Reject test",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        res = reject_reimbursement_endpoint(
            r.id,
            StoreReimbursementReject(rejection_reason="Duplicate claim detected; already compensated."),
            self.db,
            self.superadmin,
        )
        self.assertEqual(res.status, "REJECTED")
        self.assertEqual(res.rejected_by_id, self.superadmin.id)
        self.assertIn("Duplicate claim", res.rejection_reason)

    # =========================================================================
    # 5. PAYMENT LIFECYCLE TESTS (PARTIAL -> FULL -> REJECTIONS)
    # =========================================================================

    def test_cannot_pay_pending_reimbursement(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("2000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("2000.00"),
            status="PENDING",
            reason="Unapproved payment attempt",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        pay_payload = StoreReimbursementPaymentCreate(
            amount=Decimal("1000.00"),
            payment_date=datetime.now(timezone.utc),
            payment_mode="UPI",
        )
        with self.assertRaises(HTTPException) as ctx:
            record_payment_endpoint(r.id, pay_payload, self.db, self.superadmin)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("must be approved before payments can be recorded", ctx.exception.detail)

    def test_full_installment_flow_partial_then_paid(self):
        # 1. Start with APPROVED claim of 2,000
        r = StoreReimbursement(
            reimbursement_number="SR-PAY-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("2000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("2000.00"),
            status="APPROVED",
            reason="Installment payment test",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        # 2. First installment of 1,000 -> PARTIALLY_PAID
        p1 = StoreReimbursementPaymentCreate(
            amount=Decimal("1000.00"),
            payment_date=datetime.now(timezone.utc),
            payment_mode="UPI",
            transaction_reference="UPI-REF-001",
            remarks="First half payment",
        )
        res_p1 = record_payment_endpoint(r.id, p1, self.db, self.superadmin)
        self.assertEqual(res_p1.amount, Decimal("1000.00"))

        reloaded = self.db.query(StoreReimbursement).filter(StoreReimbursement.id == r.id).first()
        self.assertEqual(reloaded.status, "PARTIALLY_PAID")
        self.assertEqual(reloaded.total_paid, Decimal("1000.00"))
        self.assertEqual(reloaded.outstanding_amount, Decimal("1000.00"))

        # 3. Second installment of 500 -> STILL PARTIALLY_PAID
        p2 = StoreReimbursementPaymentCreate(
            amount=Decimal("500.00"),
            payment_date=datetime.now(timezone.utc),
            payment_mode="BANK_TRANSFER",
            transaction_reference="IMPS-REF-002",
        )
        record_payment_endpoint(r.id, p2, self.db, self.superadmin)

        self.db.refresh(reloaded)
        self.assertEqual(reloaded.status, "PARTIALLY_PAID")
        self.assertEqual(reloaded.total_paid, Decimal("1500.00"))
        self.assertEqual(reloaded.outstanding_amount, Decimal("500.00"))

        # 4. Final installment of 500 -> PAID
        p3 = StoreReimbursementPaymentCreate(
            amount=Decimal("500.00"),
            payment_date=datetime.now(timezone.utc),
            payment_mode="CASH",
        )
        record_payment_endpoint(r.id, p3, self.db, self.superadmin)

        self.db.refresh(reloaded)
        self.assertEqual(reloaded.status, "PAID")
        self.assertEqual(reloaded.total_paid, Decimal("2000.00"))
        self.assertEqual(reloaded.outstanding_amount, Decimal("0.00"))

        # 5. Further payment on PAID claim rejected
        p4 = StoreReimbursementPaymentCreate(
            amount=Decimal("100.00"),
            payment_date=datetime.now(timezone.utc),
        )
        with self.assertRaises(HTTPException) as ctx:
            record_payment_endpoint(r.id, p4, self.db, self.superadmin)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("PAID", ctx.exception.detail)

    def test_overpayment_is_strictly_rejected(self):
        r = StoreReimbursement(
            reimbursement_number="SR-OVER-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="Overpayment test",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        p = StoreReimbursementPaymentCreate(
            amount=Decimal("1500.00"),  # Exceeds 1000
            payment_date=datetime.now(timezone.utc),
        )
        with self.assertRaises(HTTPException) as ctx:
            record_payment_endpoint(r.id, p, self.db, self.superadmin)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("exceeds outstanding balance", ctx.exception.detail)

    def test_employee_cannot_record_payment(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="Emp pay check",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        p = StoreReimbursementPaymentCreate(
            amount=Decimal("500.00"),
            payment_date=datetime.now(timezone.utc),
        )
        with self.assertRaises(HTTPException) as ctx:
            record_payment_endpoint(r.id, p, self.db, self.employee)
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("Only Super Admin", ctx.exception.detail)

    def test_duplicate_transaction_reference_on_same_claim_rejected(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="Dup test",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        p1 = StoreReimbursementPaymentCreate(
            amount=Decimal("200.00"),
            payment_date=datetime.now(timezone.utc),
            transaction_reference="UNIQUE-UTR-999",
        )
        record_payment_endpoint(r.id, p1, self.db, self.superadmin)

        # Attempt to record identical reference on same claim
        p2 = StoreReimbursementPaymentCreate(
            amount=Decimal("200.00"),
            payment_date=datetime.now(timezone.utc),
            transaction_reference="UNIQUE-UTR-999",
        )
        with self.assertRaises(HTTPException) as ctx:
            record_payment_endpoint(r.id, p2, self.db, self.superadmin)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("already exists for this claim", ctx.exception.detail)

    # =========================================================================
    # 6. ATTACHMENT & VOID TESTS
    # =========================================================================

    async def test_attachment_upload_and_voiding_preserves_db_record(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Att test",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        # 1. Upload valid image attachment
        img_bytes = self._make_image_bytes()
        upload_file = UploadFile(filename="invoice.png", file=io.BytesIO(img_bytes), headers={"content-type": "image/png"})

        res_att = await upload_attachment_endpoint(
            reimbursement_id=r.id,
            file=upload_file,
            attachment_type="INVOICE",
            payment_id=None,
            db=self.db,
            current_user=self.employee,
        )
        self.assertEqual(res_att.attachment_type, "INVOICE")
        self.assertEqual(res_att.file_name, "invoice.png")
        self.assertIsNone(res_att.voided_at)

        # 2. Void attachment
        res_void = void_attachment_endpoint(
            reimbursement_id=r.id,
            attachment_id=res_att.id,
            payload=StoreReimbursementAttachmentVoidRequest(void_reason="Uploaded wrong receipt page."),
            db=self.db,
            current_user=self.employee,
        )
        self.assertIsNotNone(res_void.voided_at)
        self.assertEqual(res_void.void_reason, "Uploaded wrong receipt page.")

        # 3. Check record is STILL in DB, not hard-deleted
        att_row = self.db.query(StoreReimbursementAttachment).filter(StoreReimbursementAttachment.id == res_att.id).first()
        self.assertIsNotNone(att_row)
        self.assertIsNotNone(att_row.voided_at)

    async def test_attachment_upload_pdf_supported(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="PDF test",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        pdf_bytes = b"%PDF-1.4 sample pdf content test bytes"
        upload_file = UploadFile(filename="tax_invoice.pdf", file=io.BytesIO(pdf_bytes), headers={"content-type": "application/pdf"})

        res = await upload_attachment_endpoint(
            reimbursement_id=r.id,
            file=upload_file,
            attachment_type="INVOICE",
            payment_id=None,
            db=self.db,
            current_user=self.superadmin,
        )
        self.assertEqual(res.file_type, "application/pdf")
        self.assertEqual(res.file_name, "tax_invoice.pdf")

    async def test_attachment_linked_to_payment_of_another_reimbursement_fails(self):
        r1 = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="R1",
            created_by_id=self.superadmin.id,
        )
        r2 = StoreReimbursement(
            reimbursement_number="SR-02",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="R2",
            created_by_id=self.superadmin.id,
        )
        self.db.add_all([r1, r2])
        self.db.commit()

        # Payment on r2
        p_r2 = StoreReimbursementPayment(
            reimbursement_id=r2.id,
            amount=Decimal("500.00"),
            payment_date=datetime.now(timezone.utc),
            paid_by_id=self.superadmin.id,
        )
        self.db.add(p_r2)
        self.db.commit()

        # Attempt to upload to r1 linking p_r2
        upload_file = UploadFile(filename="invoice.png", file=io.BytesIO(self._make_image_bytes()), headers={"content-type": "image/png"})
        with self.assertRaises(HTTPException) as ctx:
            await upload_attachment_endpoint(
                reimbursement_id=r1.id,
                file=upload_file,
                attachment_type="PAYMENT_PROOF",
                payment_id=p_r2.id,
                db=self.db,
                current_user=self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("does not belong to Reimbursement", ctx.exception.detail)

    # =========================================================================
    # 7. CANCELLATION TESTS
    # =========================================================================

    def test_pending_cancellation_with_reason(self):
        r = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Cancel test",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        res = cancel_reimbursement_endpoint(
            r.id,
            StoreReimbursementCancel(cancellation_reason="Store POC resolved invoice directly with vendor."),
            self.db,
            self.employee,
        )
        self.assertEqual(res.status, "CANCELLED")
        self.assertEqual(res.cancelled_by_id, self.employee.id)
        self.assertIn("resolved invoice", res.cancellation_reason)

    def test_superadmin_cancels_approved_unpaid_claim(self):
        r = StoreReimbursement(
            reimbursement_number="SR-02",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="Unpaid approved cancel test",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        res = cancel_reimbursement_endpoint(
            r.id,
            StoreReimbursementCancel(cancellation_reason="Vendor provided direct credit note"),
            self.db,
            self.superadmin,
        )
        self.assertEqual(res.status, "CANCELLED")
        self.assertEqual(res.cancelled_by_id, self.superadmin.id)
        self.assertEqual(res.cancellation_reason, "Vendor provided direct credit note")

    def test_employee_cannot_cancel_approved_claim(self):
        r = StoreReimbursement(
            reimbursement_number="SR-03",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="APPROVED",
            reason="Approved claim test",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            cancel_reimbursement_endpoint(
                r.id,
                StoreReimbursementCancel(cancellation_reason="Attempt cancel approved"),
                self.db,
                self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("Employees can only cancel PENDING claims", ctx.exception.detail)

    def test_approved_claim_with_any_payment_cannot_be_cancelled(self):
        r = StoreReimbursement(
            reimbursement_number="SR-04",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("200.00"),
            outstanding_amount=Decimal("800.00"),
            status="APPROVED",
            reason="Approved with payment cancel attempt",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            cancel_reimbursement_endpoint(
                r.id,
                StoreReimbursementCancel(cancellation_reason="Attempt cancel with payments"),
                self.db,
                self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("Cannot cancel an approved claim with recorded payments", ctx.exception.detail)

    def test_partially_paid_claim_cannot_be_cancelled(self):
        r = StoreReimbursement(
            reimbursement_number="SR-05",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("500.00"),
            outstanding_amount=Decimal("500.00"),
            status="PARTIALLY_PAID",
            reason="Paid claim cancel attempt",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            cancel_reimbursement_endpoint(
                r.id,
                StoreReimbursementCancel(cancellation_reason="Cancel anyway"),
                self.db,
                self.superadmin,
            )
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("payments have already been executed", ctx.exception.detail)

    def test_employee_can_cancel_own_scoped_pending_claim_with_permission(self):
        r = StoreReimbursement(
            reimbursement_number="SR-06",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Own claim cancel",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        res = cancel_reimbursement_endpoint(
            r.id,
            StoreReimbursementCancel(cancellation_reason="Cancelled by creator employee"),
            self.db,
            self.employee,
        )
        self.assertEqual(res.status, "CANCELLED")
        self.assertEqual(res.cancelled_by_id, self.employee.id)

    def test_employee_cannot_cancel_another_employee_pending_claim(self):
        r = StoreReimbursement(
            reimbursement_number="SR-07",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Other employee claim",
            created_by_id=self.other_employee.id,
        )
        self.db.add(r)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            cancel_reimbursement_endpoint(
                r.id,
                StoreReimbursementCancel(cancellation_reason="Attempt hijack cancel"),
                self.db,
                self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("Employees can only cancel reimbursements they created", ctx.exception.detail)

    def test_employee_without_cancel_permission_gets_403(self):
        # Remove store_reimbursement.cancel permission from employee
        cancel_perm = self.perms["store_reimbursement.cancel"]
        self.db.query(EmployeePermission).filter_by(
            employee_id=self.employee.id,
            permission_id=cancel_perm.id,
        ).delete()
        self.db.commit()

        r = StoreReimbursement(
            reimbursement_number="SR-08",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("0.00"),
            outstanding_amount=Decimal("1000.00"),
            status="PENDING",
            reason="Cancel test without perm",
            created_by_id=self.employee.id,
        )
        self.db.add(r)
        self.db.commit()

        with self.assertRaises(HTTPException) as ctx:
            cancel_reimbursement_endpoint(
                r.id,
                StoreReimbursementCancel(cancellation_reason="Unauthorized cancel"),
                self.db,
                self.employee,
            )
        self.assertEqual(ctx.exception.status_code, 403)
        self.assertIn("permission", ctx.exception.detail.lower())

    def test_two_sequential_records_receive_distinct_stable_reimbursement_numbers(self):
        payload1 = StoreReimbursementCreate(
            store_id=self.store_1.id,
            purchase_date=date(2026, 10, 5),
            total_amount=Decimal("500.00"),
            reason="First sequential emergency purchase",
        )
        payload2 = StoreReimbursementCreate(
            store_id=self.store_1.id,
            purchase_date=date(2026, 10, 6),
            total_amount=Decimal("750.00"),
            reason="Second sequential emergency purchase",
        )

        r1 = create_reimbursement_endpoint(payload1, self.db, self.superadmin)
        r2 = create_reimbursement_endpoint(payload2, self.db, self.superadmin)

        self.assertIsNotNone(r1.reimbursement_number)
        self.assertIsNotNone(r2.reimbursement_number)
        self.assertNotEqual(r1.reimbursement_number, r2.reimbursement_number)

        import re
        self.assertTrue(re.match(r"^SR-\d{6}-\d{5}$", r1.reimbursement_number))
        self.assertTrue(re.match(r"^SR-\d{6}-\d{5}$", r2.reimbursement_number))

        id1_seq = int(r1.reimbursement_number.split("-")[-1])
        id2_seq = int(r2.reimbursement_number.split("-")[-1])
        self.assertGreater(id2_seq, id1_seq)

    # =========================================================================
    # 8. SUMMARY & EXPORT TESTS
    # =========================================================================

    def test_summary_card_calculations(self):
        r1 = StoreReimbursement(
            reimbursement_number="SR-01",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            total_paid=Decimal("500.00"),
            outstanding_amount=Decimal("500.00"),
            status="PARTIALLY_PAID",
            reason="Sum 1",
            created_by_id=self.superadmin.id,
        )
        r2 = StoreReimbursement(
            reimbursement_number="SR-02",
            store_id=self.store_1.id,
            store_name="Store 1",
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("800.00"),
            total_paid=Decimal("800.00"),
            outstanding_amount=Decimal("0.00"),
            status="PAID",
            reason="Sum 2",
            created_by_id=self.superadmin.id,
        )
        self.db.add_all([r1, r2])
        self.db.commit()

        summary = get_reimbursement_summary_endpoint(db=self.db, current_user=self.superadmin)
        self.assertEqual(summary.total_claims, 2)
        self.assertEqual(summary.partially_paid_count, 1)
        self.assertEqual(summary.paid_count, 1)
        self.assertEqual(summary.total_reimbursement_amount, Decimal("1800.00"))
        self.assertEqual(summary.total_paid_amount, Decimal("1300.00"))
        self.assertEqual(summary.total_outstanding_amount, Decimal("500.00"))

    def test_export_endpoint_generates_valid_excel_workbook(self):
        r = StoreReimbursement(
            reimbursement_number="SR-EXP-01",
            store_id=self.store_1.id,
            store_name="Gurgaon Sector 29",
            store_region="NORTH",
            store_channel="BLINKIT",
            store_entity="BCPL",
            poc_name="Partner One",
            poc_phone="9876543210",
            purchase_date=date(2026, 10, 5),
            total_amount=Decimal("1250.00"),
            total_paid=Decimal("1250.00"),
            outstanding_amount=Decimal("0.00"),
            status="PAID",
            reason="Export verify",
            created_by_id=self.superadmin.id,
        )
        self.db.add(r)
        self.db.commit()

        res = export_reimbursements_endpoint(db=self.db, current_user=self.superadmin)
        self.assertEqual(res.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("store_reimbursements_", res.headers["Content-Disposition"])

        # Validate openpyxl parse
        wb = openpyxl.load_workbook(io.BytesIO(res.body))
        ws = wb.active
        self.assertEqual(ws.title, "Store Reimbursements")
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Reimbursement No", headers)
        self.assertIn("Total Amount", headers)
        self.assertIn("Paid Amount", headers)
        self.assertIn("Outstanding Amount", headers)

        # Check row content
        row_vals = [cell.value for cell in ws[2]]
        self.assertIn("SR-EXP-01", row_vals)
        self.assertIn("Gurgaon Sector 29", row_vals)
        self.assertIn(1250.0, row_vals)

    # =========================================================================
    # 9. HISTORICAL SNAPSHOT PRESERVATION
    # =========================================================================

    def test_store_and_poc_snapshots_remain_intact_if_store_is_edited_later(self):
        payload = StoreReimbursementCreate(
            store_id=self.store_1.id,
            purchase_date=date(2026, 10, 1),
            total_amount=Decimal("1000.00"),
            reason="Snapshot integrity test",
        )
        res = create_reimbursement_endpoint(payload, self.db, self.superadmin)
        reimb_id = res.id

        # Now simulate store being renamed and partner updated
        self.store_1.store_name = "Renamed Store Name 2027"
        self.partner_1.full_name = "Changed POC Name 2027"
        self.db.commit()

        # Reload reimbursement: historical snapshots must remain unchanged!
        reloaded = self.db.query(StoreReimbursement).filter(StoreReimbursement.id == reimb_id).first()
        self.assertEqual(reloaded.store_name, "Gurgaon Sector 29")
        self.assertEqual(reloaded.poc_name, "Partner One")

    # =========================================================================
    # 10. RBAC SEEDER VERIFICATION
    # =========================================================================

    def test_permission_seeder_contains_all_store_reimbursement_permissions(self):
        rows = build_permission_rows()
        keys = {r["permission_key"] for r in rows}

        expected = [
            "store_reimbursement.view",
            "store_reimbursement.create",
            "store_reimbursement.edit",
            "store_reimbursement.approve",
            "store_reimbursement.record_payment",
            "store_reimbursement.upload_proof",
            "store_reimbursement.cancel",
            "store_reimbursement.export",
        ]
        for exp in expected:
            self.assertIn(exp, keys)
