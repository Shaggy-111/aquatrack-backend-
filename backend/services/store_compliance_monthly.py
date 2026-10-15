"""Domain service layer for Store Compliance Phase 2 (Monthly workflow, Invoices, Payments)."""

from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import io
import os
from typing import List, Optional, Tuple
import uuid

from PIL import Image, UnidentifiedImageError
from sqlalchemy import func
from sqlalchemy.orm import Session
from starlette.concurrency import run_in_threadpool

from ..models.store import Store
from ..models.user import RoleEnum, User
from ..models.store_compliance_monthly import (
    StoreComplianceAttachment,
    StoreComplianceAttachmentType,
    StoreComplianceInvoice,
    StoreCompliancePayment,
    StoreCompliancePaymentMode,
    StoreCompliancePaymentStatus,
    StoreMonthlyCardStatus,
    StoreMonthlyCompliance,
)
from ..schemas.store_compliance_monthly import (
    StoreComplianceInvoiceUpdate,
    StoreCompliancePaymentCreate,
    StoreMonthlyComplianceCreate,
)
from ..utils.permissions import get_employee_store_ids


# =============================================================================
# DOMAIN EXCEPTIONS
# =============================================================================

class StoreComplianceMonthlyDomainError(ValueError):
    """Base exception for domain rule violations."""
    pass


class StoreComplianceMonthlyNotFoundError(StoreComplianceMonthlyDomainError):
    """Resource was not found."""
    pass


class StoreComplianceMonthlyStateError(StoreComplianceMonthlyDomainError):
    """Action is invalid for the current resource lifecycle state."""
    pass


class StoreComplianceMonthlyStoreAccessError(StoreComplianceMonthlyDomainError):
    """Employee does not have access to the target store."""
    pass


class StoreComplianceMonthlyPaymentError(StoreComplianceMonthlyDomainError):
    """Payment ledger rule violation."""
    pass


class StoreComplianceMonthlyAttachmentError(StoreComplianceMonthlyDomainError):
    """Attachment upload or validation failure."""
    pass


MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024  # 10 MB
ALLOWED_MIME_TYPES = {
    "image/jpeg",
    "image/jpg",
    "image/png",
    "image/webp",
    "application/pdf",
}
ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".pdf"}


# =============================================================================
# HELPER: FILE VALIDATION & CLOUDINARY UPLOAD
# =============================================================================

async def _validate_and_upload_file(
    content: bytes,
    filename: str,
    content_type: str,
    folder_path: str,
    public_id_prefix: str,
) -> Tuple[str, Optional[str]]:
    """Validate file content and upload to Cloudinary (or mock storage)."""
    if len(content) == 0:
        raise StoreComplianceMonthlyAttachmentError("Uploaded file is empty.")

    if len(content) > MAX_ATTACHMENT_BYTES:
        raise StoreComplianceMonthlyAttachmentError("Uploaded file exceeds maximum allowed size (10 MB).")

    ext = os.path.splitext(filename)[1].lower() if filename else ""
    if ext not in ALLOWED_EXTENSIONS and content_type not in ALLOWED_MIME_TYPES:
        raise StoreComplianceMonthlyAttachmentError(
            f"Unsupported file format '{ext or content_type}'. Allowed: JPEG, PNG, WEBP, PDF."
        )

    # Content verification for image files
    if content_type.startswith("image/") or ext in {".jpg", ".jpeg", ".png", ".webp"}:
        try:
            with Image.open(io.BytesIO(content)) as img:
                img.verify()
        except (UnidentifiedImageError, OSError) as exc:
            raise StoreComplianceMonthlyAttachmentError("Invalid or corrupted image file.") from exc

    cloud_name = os.getenv("CLOUDINARY_CLOUD_NAME")
    if cloud_name:
        import cloudinary
        import cloudinary.uploader
        from ..utils.cloudinary_config import cloudinary as _c_init

        is_pdf = content_type == "application/pdf" or ext == ".pdf"
        resource_type = "raw" if is_pdf else "image"
        public_id = f"{public_id_prefix}_{uuid.uuid4().hex[:12]}"

        try:
            result = await run_in_threadpool(
                cloudinary.uploader.upload,
                io.BytesIO(content),
                resource_type=resource_type,
                type="upload",
                folder=folder_path,
                public_id=public_id,
            )
            return result["secure_url"], result.get("public_id")
        except Exception as exc:
            raise StoreComplianceMonthlyAttachmentError(f"Cloudinary upload failed: {str(exc)}") from exc
    else:
        mock_id = f"mock_{uuid.uuid4().hex[:12]}"
        mock_url = f"https://mock-storage.aquatrack.internal/compliance/{mock_id}_{filename}"
        return mock_url, mock_id


# =============================================================================
# 1. MONTHLY COMPLIANCE CREATION
# =============================================================================

def create_store_monthly_compliance(
    payload: StoreMonthlyComplianceCreate,
    db: Session,
    current_user: User,
) -> StoreMonthlyCompliance:
    """Initialize a monthly compliance record for a Store, Year, and Month."""
    if not (1 <= payload.month <= 12):
        raise StoreComplianceMonthlyDomainError("Month must be between 1 and 12.")

    if not (2000 <= payload.year <= 2100):
        raise StoreComplianceMonthlyDomainError("Year must be between 2000 and 2100.")

    # 1. Store validation
    store = db.query(Store).filter(Store.id == payload.store_id.strip()).first()
    if not store:
        raise StoreComplianceMonthlyNotFoundError(f"Store '{payload.store_id}' not found.")

    # 2. Scope check for employees
    if current_user.role == RoleEnum.employee:
        assigned_ids = get_employee_store_ids(current_user.id, db=db)
        if store.id not in assigned_ids:
            raise StoreComplianceMonthlyStoreAccessError(
                f"Store '{store.id}' is not assigned to your employee account."
            )

    # 3. Uniqueness check
    existing = db.query(StoreMonthlyCompliance).filter(
        StoreMonthlyCompliance.store_id == store.id,
        StoreMonthlyCompliance.year == payload.year,
        StoreMonthlyCompliance.month == payload.month,
    ).first()
    if existing:
        raise StoreComplianceMonthlyDomainError(
            f"Monthly compliance record already exists for Store '{store.id}', {payload.year}-{payload.month:02d}."
        )

    # 4. Snapshot delivery manager
    dm_id = store.assigned_manager_id
    dm_name = store.assigned_manager.full_name if store.assigned_manager else None

    # 5. Initialize entity with temporary unique placeholder to satisfy NOT NULL before PK flush
    temp_placeholder = f"TMP-{uuid.uuid4().hex}"
    now_utc = datetime.now(timezone.utc)

    compliance = StoreMonthlyCompliance(
        compliance_number=temp_placeholder,
        store_id=store.id,
        store_name=store.store_name,
        store_channel=store.channel,
        store_region=store.region,
        store_entity=store.entity,
        delivery_manager_id=dm_id,
        delivery_manager_name=dm_name,
        year=payload.year,
        month=payload.month,
        card_status=StoreMonthlyCardStatus.PENDING_CARD.value,
        card_remarks=payload.card_remarks.strip() if payload.card_remarks else None,
        created_by_id=current_user.id,
        created_at=now_utc,
        updated_at=now_utc,
    )

    db.add(compliance)
    db.flush()

    # 6. Final deterministic compliance identifier: SMC-{YYYYMM}-{id:05d}
    compliance.compliance_number = f"SMC-{payload.year}{payload.month:02d}-{compliance.id:05d}"
    db.flush()

    return compliance


# =============================================================================
# 2. MONTHLY WATER CARD UPLOAD & REPLACEMENT
# =============================================================================

async def upload_monthly_card(
    compliance: StoreMonthlyCompliance,
    content: bytes,
    filename: str,
    content_type: str,
    remarks: Optional[str],
    db: Session,
    current_user: User,
) -> StoreMonthlyCompliance:
    """Upload or replace the primary monthly water card scan without audit evidence loss."""
    folder = f"aquatrack/store-compliance/{compliance.store_id}/{compliance.year}/{compliance.month:02d}/cards"
    file_url, file_public_id = await _validate_and_upload_file(
        content=content,
        filename=filename,
        content_type=content_type,
        folder_path=folder,
        public_id_prefix="card",
    )

    now_utc = datetime.now(timezone.utc)

    # 1. Soft-void the existing active CARD attachment if one exists (do NOT duplicate)
    active_card_att = (
        db.query(StoreComplianceAttachment)
        .filter(
            StoreComplianceAttachment.monthly_compliance_id == compliance.id,
            StoreComplianceAttachment.attachment_type == StoreComplianceAttachmentType.CARD.value,
            StoreComplianceAttachment.voided_at.is_(None),
        )
        .first()
    )
    if active_card_att:
        active_card_att.voided_at = now_utc
        active_card_att.voided_by_id = current_user.id
        active_card_att.void_reason = "Replaced by new monthly water card upload"
    elif compliance.card_file_url:
        # Fallback if card existed prior to attachments ledger tracking
        prior_attachment = StoreComplianceAttachment(
            monthly_compliance_id=compliance.id,
            attachment_type=StoreComplianceAttachmentType.CARD.value,
            file_url=compliance.card_file_url,
            file_public_id=compliance.card_file_public_id,
            file_name=compliance.card_file_name or "prior_card_upload",
            file_type=compliance.card_file_type,
            file_size=compliance.card_file_size,
            uploaded_by_id=compliance.card_uploaded_by_id or current_user.id,
            created_at=compliance.card_uploaded_at or now_utc,
            voided_at=now_utc,
            voided_by_id=current_user.id,
            void_reason="Replaced by new monthly water card upload",
        )
        db.add(prior_attachment)

    # 2. Set new primary card file on monthly compliance
    compliance.card_status = StoreMonthlyCardStatus.CARD_UPLOADED.value
    compliance.card_file_url = file_url
    compliance.card_file_name = filename
    compliance.card_file_size = len(content)
    compliance.card_file_type = content_type
    compliance.card_file_public_id = file_public_id
    if remarks is not None:
        compliance.card_remarks = remarks.strip() if remarks else None
    compliance.card_uploaded_by_id = current_user.id
    compliance.card_uploaded_at = now_utc
    compliance.updated_by_id = current_user.id
    compliance.updated_at = now_utc

    # 3. Add active attachment entry for the new card
    active_attachment = StoreComplianceAttachment(
        monthly_compliance_id=compliance.id,
        attachment_type=StoreComplianceAttachmentType.CARD.value,
        file_url=file_url,
        file_public_id=file_public_id,
        file_name=filename,
        file_type=content_type,
        file_size=len(content),
        uploaded_by_id=current_user.id,
        created_at=now_utc,
    )
    db.add(active_attachment)
    db.flush()

    return compliance


# =============================================================================
# 3. INVOICE CREATION & UPDATES
# =============================================================================

async def create_compliance_invoice(
    compliance: StoreMonthlyCompliance,
    invoice_number: str,
    invoice_date: date,
    invoice_amount: Decimal,
    content: bytes,
    filename: str,
    content_type: str,
    remarks: Optional[str] = None,
    db: Session = None,
    current_user: User = None,
) -> StoreComplianceInvoice:
    """Attach commercial invoice with mandatory invoice document (SuperAdmin only)."""
    if current_user.role != RoleEnum.superadmin:
        raise StoreComplianceMonthlyDomainError("Only Super Admin can create compliance invoices.")

    if compliance.invoice:
        raise StoreComplianceMonthlyStateError("An invoice is already attached to this monthly compliance record.")

    inv_num_clean = invoice_number.strip() if invoice_number else ""
    if not inv_num_clean:
        raise StoreComplianceMonthlyDomainError("Invoice number is required.")

    if not invoice_date:
        raise StoreComplianceMonthlyDomainError("Invoice date is required.")

    amount = Decimal(str(invoice_amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if amount <= Decimal("0.00"):
        raise StoreComplianceMonthlyDomainError("Invoice amount must be strictly greater than 0.")

    if not content or len(content) == 0:
        raise StoreComplianceMonthlyAttachmentError("An invoice document (PDF or image) is required to create an invoice.")

    folder = f"aquatrack/store-compliance/{compliance.store_id}/{compliance.year}/{compliance.month:02d}/invoices"
    file_url, file_public_id = await _validate_and_upload_file(
        content=content,
        filename=filename or "invoice",
        content_type=content_type or "application/octet-stream",
        folder_path=folder,
        public_id_prefix="inv",
    )

    now_utc = datetime.now(timezone.utc)

    invoice = StoreComplianceInvoice(
        monthly_compliance_id=compliance.id,
        invoice_number=inv_num_clean,
        invoice_date=invoice_date,
        invoice_amount=amount,
        total_paid=Decimal("0.00"),
        outstanding_amount=amount,
        payment_status=StoreCompliancePaymentStatus.UNPAID.value,
        file_url=file_url,
        file_name=filename,
        file_size=len(content),
        file_type=content_type,
        file_public_id=file_public_id,
        remarks=remarks.strip() if remarks else None,
        created_by_id=current_user.id,
        created_at=now_utc,
        updated_at=now_utc,
    )
    db.add(invoice)
    db.flush()

    # Track primary invoice attachment in attachments table
    attachment = StoreComplianceAttachment(
        monthly_compliance_id=compliance.id,
        invoice_id=invoice.id,
        attachment_type=StoreComplianceAttachmentType.INVOICE.value,
        file_url=file_url,
        file_public_id=file_public_id,
        file_name=filename,
        file_type=content_type,
        file_size=len(content),
        uploaded_by_id=current_user.id,
        created_at=now_utc,
    )
    db.add(attachment)
    db.flush()

    compliance.updated_by_id = current_user.id
    compliance.updated_at = now_utc
    db.flush()

    return invoice


async def update_compliance_invoice(
    invoice: StoreComplianceInvoice,
    payload: StoreComplianceInvoiceUpdate,
    new_file_content: Optional[bytes] = None,
    new_filename: Optional[str] = None,
    new_content_type: Optional[str] = None,
    db: Session = None,
    current_user: User = None,
) -> StoreComplianceInvoice:
    """Edit invoice fields or replace document before payment begins (SuperAdmin only)."""
    if current_user.role != RoleEnum.superadmin:
        raise StoreComplianceMonthlyDomainError("Only Super Admin can update invoices.")

    has_payments = invoice.total_paid > Decimal("0.00")
    now_utc = datetime.now(timezone.utc)

    if has_payments:
        if (
            (payload.invoice_number is not None and payload.invoice_number.strip() != invoice.invoice_number)
            or (payload.invoice_date is not None and payload.invoice_date != invoice.invoice_date)
            or (payload.invoice_amount is not None and payload.invoice_amount != invoice.invoice_amount)
            or (new_file_content is not None and len(new_file_content) > 0)
        ):
            raise StoreComplianceMonthlyStateError(
                "Invoice number, date, amount, and document are immutable once payments have been recorded."
            )

    if not has_payments:
        if payload.invoice_number is not None:
            clean_num = payload.invoice_number.strip()
            if not clean_num:
                raise StoreComplianceMonthlyDomainError("Invoice number cannot be empty.")
            invoice.invoice_number = clean_num

        if payload.invoice_date is not None:
            invoice.invoice_date = payload.invoice_date

        if payload.invoice_amount is not None:
            new_amount = Decimal(str(payload.invoice_amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if new_amount <= Decimal("0.00"):
                raise StoreComplianceMonthlyDomainError("Invoice amount must be strictly greater than 0.")
            invoice.invoice_amount = new_amount
            invoice.outstanding_amount = new_amount

        # File replacement before payments
        if new_file_content is not None and len(new_file_content) > 0:
            compliance = invoice.monthly_compliance
            folder = f"aquatrack/store-compliance/{compliance.store_id}/{compliance.year}/{compliance.month:02d}/invoices"
            new_url, new_public_id = await _validate_and_upload_file(
                content=new_file_content,
                filename=new_filename or "invoice_replacement",
                content_type=new_content_type or "application/octet-stream",
                folder_path=folder,
                public_id_prefix="inv",
            )

            # Soft-void the previous active INVOICE attachment
            prior_att = (
                db.query(StoreComplianceAttachment)
                .filter(
                    StoreComplianceAttachment.invoice_id == invoice.id,
                    StoreComplianceAttachment.attachment_type == StoreComplianceAttachmentType.INVOICE.value,
                    StoreComplianceAttachment.voided_at.is_(None),
                )
                .first()
            )
            if prior_att:
                prior_att.voided_at = now_utc
                prior_att.voided_by_id = current_user.id
                prior_att.void_reason = "Replaced by new invoice document upload"

            # Update invoice record with new document
            invoice.file_url = new_url
            invoice.file_name = new_filename
            invoice.file_size = len(new_file_content)
            invoice.file_type = new_content_type
            invoice.file_public_id = new_public_id

            # Add new active INVOICE attachment
            new_att = StoreComplianceAttachment(
                monthly_compliance_id=compliance.id,
                invoice_id=invoice.id,
                attachment_type=StoreComplianceAttachmentType.INVOICE.value,
                file_url=new_url,
                file_public_id=new_public_id,
                file_name=new_filename or "invoice_replacement",
                file_type=new_content_type,
                file_size=len(new_file_content),
                uploaded_by_id=current_user.id,
                created_at=now_utc,
            )
            db.add(new_att)

    if payload.remarks is not None:
        invoice.remarks = payload.remarks.strip() if payload.remarks else None

    invoice.updated_by_id = current_user.id
    invoice.updated_at = now_utc
    db.flush()

    return invoice


# =============================================================================
# 4. PAYMENT INSTALLMENT LEDGER
# =============================================================================

def record_compliance_payment(
    compliance_id: int,
    payload: StoreCompliancePaymentCreate,
    db: Session,
    current_user: User,
) -> Tuple[StoreCompliancePayment, StoreComplianceInvoice]:
    """Record an append-only payment installment against a monthly compliance invoice (SuperAdmin only)."""
    if current_user.role != RoleEnum.superadmin:
        raise StoreComplianceMonthlyDomainError("Only Super Admin can record payments.")

    # Concurrency-safe row-level lock on invoice
    invoice = (
        db.query(StoreComplianceInvoice)
        .filter(StoreComplianceInvoice.monthly_compliance_id == compliance_id)
        .with_for_update()
        .first()
    )
    if not invoice:
        raise StoreComplianceMonthlyNotFoundError(
            f"No invoice found for monthly compliance ID {compliance_id}. Invoice must exist before recording payment."
        )

    if invoice.payment_status == StoreCompliancePaymentStatus.PAID.value:
        raise StoreComplianceMonthlyPaymentError("Invoice is already fully PAID. No further payments can be recorded.")

    # Amount validation
    amount = Decimal(str(payload.amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if amount <= Decimal("0.00"):
        raise StoreComplianceMonthlyPaymentError("Payment amount must be strictly greater than 0.")

    outstanding = Decimal(str(invoice.outstanding_amount)).quantize(Decimal("0.01"))
    if amount > outstanding:
        raise StoreComplianceMonthlyPaymentError(
            f"Payment amount ({amount}) exceeds outstanding balance ({outstanding})."
        )

    # Payment Mode Validation
    valid_modes = {m.value for m in StoreCompliancePaymentMode}
    mode_clean = (payload.payment_mode or "UPI").strip().upper()
    if mode_clean not in valid_modes:
        raise StoreComplianceMonthlyPaymentError(
            f"Invalid payment mode '{mode_clean}'. Allowed: {', '.join(sorted(valid_modes))}."
        )

    # Duplicate payment safeguard on same invoice
    if payload.transaction_reference and payload.transaction_reference.strip():
        ref_clean = payload.transaction_reference.strip()
        existing_dup = (
            db.query(StoreCompliancePayment)
            .filter(
                StoreCompliancePayment.invoice_id == invoice.id,
                StoreCompliancePayment.transaction_reference == ref_clean,
                StoreCompliancePayment.amount == amount,
            )
            .first()
        )
        if existing_dup:
            raise StoreComplianceMonthlyPaymentError(
                f"A payment with transaction reference '{ref_clean}' and amount '{amount}' already exists for this invoice."
            )
    else:
        ref_clean = None

    # Concurrency-safe, sequence-safe deterministic identifier derived from database primary key
    temp_placeholder = f"TMP-{uuid.uuid4().hex}"
    now_utc = datetime.now(timezone.utc)
    payment = StoreCompliancePayment(
        invoice_id=invoice.id,
        payment_number=temp_placeholder,
        amount=amount,
        payment_date=payload.payment_date,
        payment_mode=mode_clean,
        transaction_reference=ref_clean,
        remarks=payload.remarks.strip() if payload.remarks else None,
        paid_by_id=current_user.id,
        created_at=now_utc,
    )
    db.add(payment)
    db.flush()

    # Assign final sequence-safe unique payment identifier from allocated PK
    payment.payment_number = f"CPAY-{payment.id:07d}"
    db.flush()

    # Update invoice balances atomically
    invoice.total_paid = (invoice.total_paid + amount).quantize(Decimal("0.01"))
    invoice.outstanding_amount = (invoice.invoice_amount - invoice.total_paid).quantize(Decimal("0.01"))

    if invoice.outstanding_amount == Decimal("0.00"):
        invoice.payment_status = StoreCompliancePaymentStatus.PAID.value
    else:
        invoice.payment_status = StoreCompliancePaymentStatus.PARTIALLY_PAID.value

    invoice.updated_by_id = current_user.id
    invoice.updated_at = now_utc

    # Update parent compliance timestamp
    if invoice.monthly_compliance:
        invoice.monthly_compliance.updated_by_id = current_user.id
        invoice.monthly_compliance.updated_at = now_utc

    db.flush()

    return payment, invoice


# =============================================================================
# 5. ATTACHMENT UPLOAD & VOIDING
# =============================================================================

async def upload_compliance_attachment(
    compliance: StoreMonthlyCompliance,
    attachment_type: str,
    invoice_id: Optional[int],
    payment_id: Optional[int],
    content: bytes,
    filename: str,
    content_type: str,
    db: Session,
    current_user: User,
) -> StoreComplianceAttachment:
    """Upload supporting document or proof with strict relational integrity."""
    att_type_clean = (attachment_type or "OTHER").strip().upper()
    valid_types = {t.value for t in StoreComplianceAttachmentType}
    if att_type_clean not in valid_types:
        raise StoreComplianceMonthlyAttachmentError(
            f"Invalid attachment type '{att_type_clean}'. Allowed: {', '.join(sorted(valid_types))}."
        )

    # Relational ownership validation: invoice_id must belong to this compliance
    if invoice_id is not None:
        if not compliance.invoice or compliance.invoice.id != invoice_id:
            raise StoreComplianceMonthlyAttachmentError(
                f"Invoice ID {invoice_id} does not belong to monthly compliance record {compliance.id}."
            )

    # Relational ownership validation: payment_id must belong to the invoice of this compliance
    if payment_id is not None:
        payment = db.query(StoreCompliancePayment).filter(StoreCompliancePayment.id == payment_id).first()
        if not payment:
            raise StoreComplianceMonthlyNotFoundError(f"Payment with ID {payment_id} not found.")
        if not compliance.invoice or payment.invoice_id != compliance.invoice.id:
            raise StoreComplianceMonthlyAttachmentError(
                f"Payment ID {payment_id} does not belong to monthly compliance record {compliance.id}."
            )

    # Strict role restrictions on attachment types
    if current_user.role == RoleEnum.employee and att_type_clean != StoreComplianceAttachmentType.OTHER.value:
        raise StoreComplianceMonthlyDomainError(
            f"Employees cannot upload attachment type '{att_type_clean}'. Employees may only upload OTHER supporting attachments or water cards via the /card route."
        )

    if att_type_clean in {StoreComplianceAttachmentType.INVOICE.value, StoreComplianceAttachmentType.PAYMENT_PROOF.value}:
        if current_user.role != RoleEnum.superadmin:
            raise StoreComplianceMonthlyDomainError(
                "Only Super Admin can upload invoice attachments or payment proofs."
            )

    folder = f"aquatrack/store-compliance/{compliance.store_id}/{compliance.year}/{compliance.month:02d}/attachments"
    file_url, file_public_id = await _validate_and_upload_file(
        content=content,
        filename=filename,
        content_type=content_type,
        folder_path=folder,
        public_id_prefix=att_type_clean.lower(),
    )

    now_utc = datetime.now(timezone.utc)
    attachment = StoreComplianceAttachment(
        monthly_compliance_id=compliance.id,
        invoice_id=invoice_id,
        payment_id=payment_id,
        attachment_type=att_type_clean,
        file_url=file_url,
        file_public_id=file_public_id,
        file_name=filename,
        file_type=content_type,
        file_size=len(content),
        uploaded_by_id=current_user.id,
        created_at=now_utc,
    )
    db.add(attachment)
    db.flush()

    return attachment


def void_compliance_attachment(
    compliance_id: int,
    attachment_id: int,
    void_reason: str,
    db: Session,
    current_user: User,
) -> StoreComplianceAttachment:
    """Soft-void a supporting document without deleting audit evidence (SuperAdmin only)."""
    if current_user.role != RoleEnum.superadmin:
        raise StoreComplianceMonthlyDomainError("Only Super Admin can void attachments.")

    reason_clean = void_reason.strip() if void_reason else ""
    if len(reason_clean) < 3:
        raise StoreComplianceMonthlyDomainError("A valid void reason (at least 3 characters) is required.")

    attachment = (
        db.query(StoreComplianceAttachment)
        .filter(
            StoreComplianceAttachment.id == attachment_id,
            StoreComplianceAttachment.monthly_compliance_id == compliance_id,
        )
        .first()
    )
    if not attachment:
        raise StoreComplianceMonthlyNotFoundError(
            f"Attachment ID {attachment_id} not found on monthly compliance record {compliance_id}."
        )

    if attachment.voided_at is not None:
        raise StoreComplianceMonthlyStateError("Attachment has already been voided.")

    attachment.voided_at = datetime.now(timezone.utc)
    attachment.voided_by_id = current_user.id
    attachment.void_reason = reason_clean
    db.flush()

    return attachment


# =============================================================================
# 6. EXCEL WORKBOOK GENERATOR
# =============================================================================

def generate_compliance_monthly_excel(
    rows: List[StoreMonthlyCompliance],
) -> io.BytesIO:
    """Generate professional Excel workbook (.xlsx) for monthly compliance records."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Store Compliance Monthly"
    ws.views.sheetView[0].showGridLines = True

    # Styling constants
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9)
    bold_font = Font(name="Arial", size=9, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    headers = [
        "Compliance No",
        "Store ID",
        "Store Name",
        "Region",
        "Channel",
        "Entity",
        "Month",
        "Year",
        "Card Status",
        "Invoice Number",
        "Invoice Date",
        "Invoice Amount",
        "Payment Status",
        "Total Paid",
        "Outstanding",
        "Created At",
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[1].height = 24

    total_inv_sum = Decimal("0.00")
    total_paid_sum = Decimal("0.00")
    total_out_sum = Decimal("0.00")

    for row_idx, r in enumerate(rows, start=2):
        inv = r.invoice
        inv_num = inv.invoice_number if inv else "-"
        inv_dt = inv.invoice_date.strftime("%Y-%m-%d") if inv else "-"
        inv_amt = float(inv.invoice_amount) if inv else 0.0
        pay_status = inv.payment_status if inv else "NO_INVOICE"
        paid_amt = float(inv.total_paid) if inv else 0.0
        out_amt = float(inv.outstanding_amount) if inv else 0.0

        if inv:
            total_inv_sum += inv.invoice_amount
            total_paid_sum += inv.total_paid
            total_out_sum += inv.outstanding_amount

        created_str = r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "-"

        row_data = [
            (r.compliance_number, align_center),
            (r.store_id, align_center),
            (r.store_name, align_left),
            (r.store_region or "-", align_center),
            (r.store_channel or "-", align_center),
            (r.store_entity or "-", align_center),
            (r.month, align_center),
            (r.year, align_center),
            (r.card_status, align_center),
            (inv_num, align_center),
            (inv_dt, align_center),
            (inv_amt, align_right),
            (pay_status, align_center),
            (paid_amt, align_right),
            (out_amt, align_right),
            (created_str, align_center),
        ]

        for col_idx, (val, alignment) in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = alignment
            cell.border = thin_border
            if col_idx in {12, 14, 15}:
                cell.number_format = "#,##0.00"

        ws.row_dimensions[row_idx].height = 18

    # Summary Row
    summary_row = len(rows) + 2
    ws.cell(row=summary_row, column=3, value="Total").font = bold_font
    ws.cell(row=summary_row, column=12, value=float(total_inv_sum)).font = bold_font
    ws.cell(row=summary_row, column=12).number_format = "#,##0.00"
    ws.cell(row=summary_row, column=14, value=float(total_paid_sum)).font = bold_font
    ws.cell(row=summary_row, column=14).number_format = "#,##0.00"
    ws.cell(row=summary_row, column=15, value=float(total_out_sum)).font = bold_font
    ws.cell(row=summary_row, column=15).number_format = "#,##0.00"

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
