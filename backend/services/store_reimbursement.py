import io
import os
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError
from sqlalchemy import func
from sqlalchemy.orm import Query, Session
from starlette.concurrency import run_in_threadpool
from fastapi import UploadFile

from ..models.partner_store import PartnerStore
from ..models.store import Store
from ..models.store_reimbursement import (
    StoreReimbursement,
    StoreReimbursementAttachment,
    StoreReimbursementAttachmentType,
    StoreReimbursementPayment,
    StoreReimbursementStatus,
)
from ..models.user import RoleEnum, User
from ..schemas.store_reimbursement import (
    StoreReimbursementCreate,
    StoreReimbursementPaymentCreate,
    StoreReimbursementUpdate,
)
from ..utils.permissions import get_employee_store_ids


MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024  # 10 MB
ALLOWED_ATTACHMENT_TYPES = {
    "image/jpeg",
    "image/jpg",
    "image/png",
    "image/webp",
    "application/pdf",
}

INTERNAL_ROLES = {
    RoleEnum.employee,
    RoleEnum.admin,
    RoleEnum.superadmin,
    RoleEnum.delivery_manager,
    RoleEnum.channel_admin,
}


# =============================================================================
# DOMAIN EXCEPTIONS
# =============================================================================

class StoreReimbursementDomainError(ValueError):
    """Base domain exception for store reimbursements."""


class StoreReimbursementNotFoundError(StoreReimbursementDomainError):
    pass


class StoreReimbursementStoreAccessError(StoreReimbursementDomainError):
    pass


class StoreReimbursementStateError(StoreReimbursementDomainError):
    pass


class StoreReimbursementPaymentError(StoreReimbursementDomainError):
    pass


class StoreReimbursementPOCError(StoreReimbursementDomainError):
    pass


class StoreReimbursementAttachmentError(StoreReimbursementDomainError):
    pass


# =============================================================================
# BUSINESS LOGIC SERVICE
# =============================================================================

def create_store_reimbursement(
    payload: StoreReimbursementCreate,
    db: Session,
    current_user: User,
) -> StoreReimbursement:
    # 1. Resolve Store
    store = db.query(Store).filter(Store.id == payload.store_id).first()
    if not store:
        raise StoreReimbursementNotFoundError(f"Store with ID '{payload.store_id}' does not exist.")

    # 2. Scope check for Employees
    if current_user.role == RoleEnum.employee:
        assigned_ids = get_employee_store_ids(current_user.id, db=db)
        if store.id not in assigned_ids:
            raise StoreReimbursementStoreAccessError(
                f"Store '{store.id}' is not assigned to your employee account."
            )

    # 3. Resolve POC
    poc_user: Optional[User] = None
    if payload.poc_id is not None:
        partner = db.query(User).filter(
            User.id == payload.poc_id,
            User.role == RoleEnum.partner,
        ).first()
        if not partner:
            raise StoreReimbursementPOCError(f"User with ID {payload.poc_id} is not a valid Partner/POC.")

        # Verify partner is mapped to this store
        mapping = db.query(PartnerStore).filter(
            PartnerStore.user_id == partner.id,
            PartnerStore.store_id == store.id,
        ).first()
        if not mapping:
            raise StoreReimbursementPOCError(
                f"Partner '{partner.full_name}' is not legitimately associated with Store '{store.id}'."
            )
        poc_user = partner
    else:
        # Check partners associated with store
        store_partners = db.query(User).join(
            PartnerStore, PartnerStore.user_id == User.id
        ).filter(
            PartnerStore.store_id == store.id,
            User.role == RoleEnum.partner,
        ).all()

        if len(store_partners) == 1:
            poc_user = store_partners[0]
        elif len(store_partners) > 1:
            raise StoreReimbursementPOCError(
                f"Multiple POCs exist for Store '{store.id}'. Please select an explicit payee POC."
            )
        else:
            # Check virtual property fallback
            if store.partners and len(store.partners) == 1:
                poc_user = store.partners[0]
            elif store.poc_name or store.poc_phone:
                poc_user = None  # Will record snapshot from store properties below
            else:
                raise StoreReimbursementPOCError(
                    f"Store '{store.id}' has no configured POC. Cannot create reimbursement without a payee."
                )

    poc_id_val = poc_user.id if poc_user else None
    poc_name_val = poc_user.full_name if poc_user else store.poc_name
    poc_phone_val = poc_user.mobile_number if poc_user else store.poc_phone

    # 4. Resolve "Instructed By" identity
    instructed_by_user: Optional[User] = None
    if payload.instructed_by_user_id is not None:
        inst_user = db.query(User).filter(User.id == payload.instructed_by_user_id).first()
        if not inst_user:
            raise StoreReimbursementDomainError(
                f"User with ID {payload.instructed_by_user_id} does not exist."
            )
        if inst_user.role not in INTERNAL_ROLES:
            raise StoreReimbursementDomainError(
                f"User '{inst_user.full_name}' has role '{inst_user.role.value}', which is not an internal operational role."
            )
        instructed_by_user = inst_user

    # 5. Financial amounts
    total = Decimal(str(payload.total_amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if total <= Decimal("0.00"):
        raise StoreReimbursementDomainError("Reimbursement amount must be strictly greater than 0.")

    # 6. Initialize Reimbursement entity
    temp_num = f"TMP-{uuid.uuid4().hex}"
    reimbursement = StoreReimbursement(
        reimbursement_number=temp_num,
        store_id=store.id,
        store_name=store.store_name,
        store_channel=store.channel,
        store_region=store.region,
        store_entity=store.entity,
        poc_id=poc_id_val,
        poc_name=poc_name_val,
        poc_phone=poc_phone_val,
        purchase_date=payload.purchase_date,
        total_amount=total,
        total_paid=Decimal("0.00"),
        outstanding_amount=total,
        status=StoreReimbursementStatus.PENDING.value,
        reason=payload.reason.strip(),
        notes=payload.notes.strip() if payload.notes else None,
        instructed_by_user_id=instructed_by_user.id if instructed_by_user else None,
        instructed_by_name=instructed_by_user.full_name if instructed_by_user else None,
        created_by_id=current_user.id,
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
    )

    db.add(reimbursement)
    db.flush()

    # 7. Generate deterministic, sequential, unique reimbursement number
    created_dt = reimbursement.created_at or datetime.now(timezone.utc)
    reimbursement.reimbursement_number = f"SR-{created_dt.strftime('%Y%m')}-{reimbursement.id:05d}"
    db.flush()

    return reimbursement


def update_store_reimbursement(
    reimbursement: StoreReimbursement,
    payload: StoreReimbursementUpdate,
    db: Session,
    current_user: User,
) -> StoreReimbursement:
    # 1. State check: Only PENDING allowed
    if reimbursement.status != StoreReimbursementStatus.PENDING.value:
        raise StoreReimbursementStateError(
            f"Cannot edit reimbursement in '{reimbursement.status}' status. Only PENDING claims can be edited."
        )

    # 2. Scope check for Employees
    if current_user.role == RoleEnum.employee:
        assigned_ids = get_employee_store_ids(current_user.id, db=db)
        if reimbursement.store_id not in assigned_ids:
            raise StoreReimbursementStoreAccessError(
                f"Store '{reimbursement.store_id}' is not assigned to your employee account."
            )

    # 3. Apply updates
    if payload.purchase_date is not None:
        reimbursement.purchase_date = payload.purchase_date

    if payload.reason is not None:
        reimbursement.reason = payload.reason.strip()

    if payload.notes is not None:
        reimbursement.notes = payload.notes.strip()

    if payload.total_amount is not None:
        new_total = Decimal(str(payload.total_amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if new_total <= Decimal("0.00"):
            raise StoreReimbursementDomainError("Reimbursement amount must be strictly greater than 0.")
        reimbursement.total_amount = new_total
        reimbursement.outstanding_amount = new_total  # total_paid is 0 in PENDING

    if payload.poc_id is not None:
        partner = db.query(User).filter(
            User.id == payload.poc_id,
            User.role == RoleEnum.partner,
        ).first()
        if not partner:
            raise StoreReimbursementPOCError(f"User with ID {payload.poc_id} is not a valid Partner/POC.")

        mapping = db.query(PartnerStore).filter(
            PartnerStore.user_id == partner.id,
            PartnerStore.store_id == reimbursement.store_id,
        ).first()
        if not mapping:
            raise StoreReimbursementPOCError(
                f"Partner '{partner.full_name}' is not legitimately associated with Store '{reimbursement.store_id}'."
            )
        reimbursement.poc_id = partner.id
        reimbursement.poc_name = partner.full_name
        reimbursement.poc_phone = partner.mobile_number

    if payload.instructed_by_user_id is not None:
        inst_user = db.query(User).filter(User.id == payload.instructed_by_user_id).first()
        if not inst_user or inst_user.role not in INTERNAL_ROLES:
            raise StoreReimbursementDomainError(
                "Selected instructed_by user is not a valid internal operational employee."
            )
        reimbursement.instructed_by_user_id = inst_user.id
        reimbursement.instructed_by_name = inst_user.full_name

    reimbursement.updated_by_id = current_user.id
    reimbursement.updated_at = datetime.now(timezone.utc)
    db.flush()

    return reimbursement


def approve_store_reimbursement(
    reimbursement: StoreReimbursement,
    db: Session,
    current_user: User,
) -> StoreReimbursement:
    if current_user.role != RoleEnum.superadmin:
        raise StoreReimbursementDomainError("Only Super Admin can approve reimbursements.")

    if reimbursement.status != StoreReimbursementStatus.PENDING.value:
        raise StoreReimbursementStateError(
            f"Cannot approve reimbursement in '{reimbursement.status}' status. Only PENDING claims can be approved."
        )

    reimbursement.status = StoreReimbursementStatus.APPROVED.value
    reimbursement.approved_by_id = current_user.id
    reimbursement.approved_at = datetime.now(timezone.utc)
    reimbursement.updated_by_id = current_user.id
    reimbursement.updated_at = datetime.now(timezone.utc)
    db.flush()

    return reimbursement


def reject_store_reimbursement(
    reimbursement: StoreReimbursement,
    rejection_reason: str,
    db: Session,
    current_user: User,
) -> StoreReimbursement:
    if current_user.role != RoleEnum.superadmin:
        raise StoreReimbursementDomainError("Only Super Admin can reject reimbursements.")

    if reimbursement.status != StoreReimbursementStatus.PENDING.value:
        raise StoreReimbursementStateError(
            f"Cannot reject reimbursement in '{reimbursement.status}' status. Only PENDING claims can be rejected."
        )

    reason_clean = rejection_reason.strip() if rejection_reason else ""
    if len(reason_clean) < 3:
        raise StoreReimbursementDomainError("A valid rejection reason (at least 3 characters) is required.")

    reimbursement.status = StoreReimbursementStatus.REJECTED.value
    reimbursement.rejected_by_id = current_user.id
    reimbursement.rejected_at = datetime.now(timezone.utc)
    reimbursement.rejection_reason = reason_clean
    reimbursement.updated_by_id = current_user.id
    reimbursement.updated_at = datetime.now(timezone.utc)
    db.flush()

    return reimbursement


def cancel_store_reimbursement(
    reimbursement: StoreReimbursement,
    cancellation_reason: str,
    db: Session,
    current_user: User,
) -> StoreReimbursement:
    # 1. Validate status eligibility
    if reimbursement.status in [StoreReimbursementStatus.PARTIALLY_PAID.value, StoreReimbursementStatus.PAID.value]:
        raise StoreReimbursementStateError(
            f"Cannot cancel reimbursement in '{reimbursement.status}' status because payments have already been executed."
        )

    if reimbursement.status == StoreReimbursementStatus.REJECTED.value:
        raise StoreReimbursementStateError("Cannot cancel a rejected reimbursement.")

    if reimbursement.status == StoreReimbursementStatus.CANCELLED.value:
        raise StoreReimbursementStateError("Reimbursement is already cancelled.")

    if reimbursement.status not in [StoreReimbursementStatus.PENDING.value, StoreReimbursementStatus.APPROVED.value]:
        raise StoreReimbursementStateError(f"Cannot cancel reimbursement in '{reimbursement.status}' status.")

    # 2. Authorization rules
    if current_user.role == RoleEnum.employee:
        # Employee can cancel only their own claim while still PENDING
        if reimbursement.status != StoreReimbursementStatus.PENDING.value:
            raise StoreReimbursementStateError("Employees can only cancel PENDING claims. Approved claims cannot be cancelled by employees.")
        if reimbursement.created_by_id != current_user.id:
            raise StoreReimbursementStoreAccessError("Employees can only cancel reimbursements they created.")
        assigned_ids = get_employee_store_ids(current_user.id, db=db)
        if reimbursement.store_id not in assigned_ids:
            raise StoreReimbursementStoreAccessError(
                f"Store '{reimbursement.store_id}' is not assigned to your employee account."
            )
    elif current_user.role == RoleEnum.superadmin:
        # SuperAdmin: may cancel PENDING, or APPROVED only when total_paid == 0
        if reimbursement.status == StoreReimbursementStatus.APPROVED.value:
            paid = Decimal(str(reimbursement.total_paid or 0)).quantize(Decimal("0.01"))
            if paid > Decimal("0.00"):
                raise StoreReimbursementStateError("Cannot cancel an approved claim with recorded payments.")
    else:
        raise StoreReimbursementDomainError("Only Super Admin or claim creator can cancel reimbursements.")

    reason_clean = cancellation_reason.strip() if cancellation_reason else ""
    if len(reason_clean) < 3:
        raise StoreReimbursementDomainError("A valid cancellation reason (at least 3 characters) is required.")

    reimbursement.status = StoreReimbursementStatus.CANCELLED.value
    reimbursement.cancelled_by_id = current_user.id
    reimbursement.cancelled_at = datetime.now(timezone.utc)
    reimbursement.cancellation_reason = reason_clean
    reimbursement.updated_by_id = current_user.id
    reimbursement.updated_at = datetime.now(timezone.utc)
    db.flush()

    return reimbursement


def record_reimbursement_payment(
    reimbursement_id: int,
    payload: StoreReimbursementPaymentCreate,
    db: Session,
    current_user: User,
) -> Tuple[StoreReimbursementPayment, StoreReimbursement]:
    # 1. SuperAdmin only
    if current_user.role != RoleEnum.superadmin:
        raise StoreReimbursementDomainError("Only Super Admin can record payments.")

    # 2. Concurrency-safe row-level lock
    # Use with_for_update() where supported (PostgreSQL). SQLite passes through cleanly.
    reimbursement = (
        db.query(StoreReimbursement)
        .filter(StoreReimbursement.id == reimbursement_id)
        .with_for_update()
        .first()
    )
    if not reimbursement:
        raise StoreReimbursementNotFoundError(f"Reimbursement with ID {reimbursement_id} not found.")

    # 3. Status invariant: Must be APPROVED or PARTIALLY_PAID
    if reimbursement.status == StoreReimbursementStatus.PENDING.value:
        raise StoreReimbursementPaymentError(
            "Reimbursement must be approved before payments can be recorded."
        )
    if reimbursement.status in [
        StoreReimbursementStatus.PAID.value,
        StoreReimbursementStatus.REJECTED.value,
        StoreReimbursementStatus.CANCELLED.value,
    ]:
        raise StoreReimbursementPaymentError(
            f"Cannot record payment on a '{reimbursement.status}' reimbursement."
        )

    # 4. Amount validation
    amount = Decimal(str(payload.amount)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if amount <= Decimal("0.00"):
        raise StoreReimbursementPaymentError("Payment amount must be strictly greater than 0.")

    outstanding = Decimal(str(reimbursement.outstanding_amount)).quantize(Decimal("0.01"))
    if amount > outstanding:
        raise StoreReimbursementPaymentError(
            f"Payment amount ({amount}) exceeds outstanding balance ({outstanding})."
        )

    # 5. Duplicate payment safeguard on same reimbursement
    if payload.transaction_reference and payload.transaction_reference.strip():
        ref_clean = payload.transaction_reference.strip()
        existing_dup = db.query(StoreReimbursementPayment).filter(
            StoreReimbursementPayment.reimbursement_id == reimbursement.id,
            StoreReimbursementPayment.transaction_reference == ref_clean,
            StoreReimbursementPayment.amount == amount,
        ).first()
        if existing_dup:
            raise StoreReimbursementPaymentError(
                f"A payment with transaction reference '{ref_clean}' and amount '{amount}' already exists for this claim."
            )

    # 6. Generate payment installment number
    existing_payments_count = db.query(func.count(StoreReimbursementPayment.id)).filter(
        StoreReimbursementPayment.reimbursement_id == reimbursement.id
    ).scalar() or 0

    pay_num = f"PAY-{reimbursement.id}-{existing_payments_count + 1:02d}"

    payment = StoreReimbursementPayment(
        reimbursement_id=reimbursement.id,
        payment_number=pay_num,
        amount=amount,
        payment_date=payload.payment_date,
        payment_mode=payload.payment_mode or "UPI",
        transaction_reference=payload.transaction_reference.strip() if payload.transaction_reference else None,
        remarks=payload.remarks.strip() if payload.remarks else None,
        paid_to_name=payload.paid_to_name or reimbursement.poc_name,
        paid_to_phone=payload.paid_to_phone or reimbursement.poc_phone,
        paid_by_id=current_user.id,
        created_at=datetime.now(timezone.utc),
    )
    db.add(payment)
    db.flush()

    # 7. Recalculate parent balances and state atomically
    new_total_paid = (Decimal(str(reimbursement.total_paid)) + amount).quantize(Decimal("0.01"))
    new_outstanding = (Decimal(str(reimbursement.total_amount)) - new_total_paid).quantize(Decimal("0.01"))

    reimbursement.total_paid = new_total_paid
    reimbursement.outstanding_amount = new_outstanding

    if new_outstanding == Decimal("0.00"):
        reimbursement.status = StoreReimbursementStatus.PAID.value
    else:
        reimbursement.status = StoreReimbursementStatus.PARTIALLY_PAID.value

    reimbursement.updated_by_id = current_user.id
    reimbursement.updated_at = datetime.now(timezone.utc)
    db.flush()

    return payment, reimbursement


async def upload_reimbursement_attachment(
    reimbursement_id: int,
    file: UploadFile,
    attachment_type: str,
    payment_id: Optional[int],
    db: Session,
    current_user: User,
) -> StoreReimbursementAttachment:
    reimbursement = db.query(StoreReimbursement).filter(StoreReimbursement.id == reimbursement_id).first()
    if not reimbursement:
        raise StoreReimbursementNotFoundError(f"Reimbursement with ID {reimbursement_id} not found.")

    # 1. Employee store scope check
    if current_user.role == RoleEnum.employee:
        assigned_ids = get_employee_store_ids(current_user.id, db=db)
        if reimbursement.store_id not in assigned_ids:
            raise StoreReimbursementStoreAccessError(
                f"Store '{reimbursement.store_id}' is not assigned to your employee account."
            )

    # 2. Validate payment link if provided
    payment: Optional[StoreReimbursementPayment] = None
    if payment_id is not None:
        payment = db.query(StoreReimbursementPayment).filter(
            StoreReimbursementPayment.id == payment_id,
            StoreReimbursementPayment.reimbursement_id == reimbursement.id,
        ).first()
        if not payment:
            raise StoreReimbursementAttachmentError(
                f"Payment ID {payment_id} does not belong to Reimbursement {reimbursement_id}."
            )

    # 3. Validate file type and extension
    content_type = (file.content_type or "").lower()
    filename = file.filename or "attachment"
    ext = os.path.splitext(filename)[1].lower()

    if content_type not in ALLOWED_ATTACHMENT_TYPES:
        if ext in [".jpg", ".jpeg"]:
            content_type = "image/jpeg"
        elif ext == ".png":
            content_type = "image/png"
        elif ext == ".webp":
            content_type = "image/webp"
        elif ext == ".pdf":
            content_type = "application/pdf"
        else:
            raise StoreReimbursementAttachmentError(
                "Unsupported file type. Only JPEG, PNG, WebP, and PDF files are allowed."
            )

    # 4. Validate file size (10 MB max)
    content = await file.read(MAX_ATTACHMENT_BYTES + 1)
    await file.close()

    if len(content) > MAX_ATTACHMENT_BYTES:
        raise StoreReimbursementAttachmentError("Uploaded file exceeds maximum allowed size (10 MB).")

    if len(content) == 0:
        raise StoreReimbursementAttachmentError("Uploaded file is empty.")

    # 5. Content verification (for images)
    if content_type.startswith("image/"):
        try:
            with Image.open(io.BytesIO(content)) as img:
                img.verify()
        except (UnidentifiedImageError, OSError) as exc:
            raise StoreReimbursementAttachmentError("Invalid or corrupted image file.") from exc

    # 6. Upload to Cloudinary (or mock fallback for offline/test environments)
    cloud_name = os.getenv("CLOUDINARY_CLOUD_NAME")
    if cloud_name:
        import cloudinary
        import cloudinary.uploader
        from ..utils.cloudinary_config import cloudinary as _c_init  # Ensure configured

        is_pdf = content_type == "application/pdf" or ext == ".pdf"
        resource_type = "raw" if is_pdf else "image"
        p_date = reimbursement.purchase_date or date.today()
        folder = f"aquatrack/store-reimbursements/{reimbursement.store_id}/{p_date.year}/{p_date.month:02d}"
        public_id = f"reimb_{uuid.uuid4().hex[:12]}"

        try:
            result = await run_in_threadpool(
                cloudinary.uploader.upload,
                io.BytesIO(content),
                resource_type=resource_type,
                type="upload",
                folder=folder,
                public_id=public_id,
            )
            file_url = result["secure_url"]
            file_public_id = result.get("public_id")
        except Exception as exc:
            raise StoreReimbursementAttachmentError(f"Cloudinary upload failed: {str(exc)}") from exc
    else:
        # Offline/Testing mock storage
        file_public_id = f"mock_{uuid.uuid4().hex[:12]}"
        file_url = f"https://mock-storage.aquatrack.internal/reimbursements/{file_public_id}_{filename}"

    # 7. Normalize attachment type
    att_type_clean = (attachment_type or "INVOICE").strip().upper()
    if att_type_clean not in {"INVOICE", "PAYMENT_PROOF", "OTHER"}:
        att_type_clean = "OTHER"

    attachment = StoreReimbursementAttachment(
        reimbursement_id=reimbursement.id,
        payment_id=payment.id if payment else None,
        attachment_type=att_type_clean,
        file_url=file_url,
        file_public_id=file_public_id,
        file_name=filename,
        file_type=content_type,
        file_size=len(content),
        uploaded_by_id=current_user.id,
        created_at=datetime.now(timezone.utc),
    )

    db.add(attachment)
    db.flush()

    return attachment


def void_reimbursement_attachment(
    reimbursement_id: int,
    attachment_id: int,
    void_reason: str,
    db: Session,
    current_user: User,
) -> StoreReimbursementAttachment:
    attachment = db.query(StoreReimbursementAttachment).filter(
        StoreReimbursementAttachment.id == attachment_id,
        StoreReimbursementAttachment.reimbursement_id == reimbursement_id,
    ).first()
    if not attachment:
        raise StoreReimbursementNotFoundError(f"Attachment with ID {attachment_id} not found.")

    if attachment.voided_at is not None:
        raise StoreReimbursementAttachmentError("Attachment is already voided.")

    reimbursement = attachment.reimbursement

    # Employee scope check & rules
    if current_user.role == RoleEnum.employee:
        assigned_ids = get_employee_store_ids(current_user.id, db=db)
        if reimbursement.store_id not in assigned_ids:
            raise StoreReimbursementStoreAccessError(
                f"Store '{reimbursement.store_id}' is not assigned to your employee account."
            )
        # Employee can only void their own uploaded attachment and only while PENDING
        if attachment.uploaded_by_id != current_user.id:
            raise StoreReimbursementStoreAccessError("Employees can only void attachments they personally uploaded.")
        if reimbursement.status != StoreReimbursementStatus.PENDING.value:
            raise StoreReimbursementStateError("Employees can only void attachments on PENDING claims.")
    elif current_user.role != RoleEnum.superadmin:
        raise StoreReimbursementDomainError("Only Super Admin or uploader can void attachments.")

    reason_clean = void_reason.strip() if void_reason else ""
    if len(reason_clean) < 3:
        raise StoreReimbursementDomainError("A valid void reason (at least 3 characters) is required.")

    attachment.voided_at = datetime.now(timezone.utc)
    attachment.voided_by_id = current_user.id
    attachment.void_reason = reason_clean
    db.flush()

    return attachment


# =============================================================================
# EXCEL EXPORT GENERATOR
# =============================================================================

def generate_reimbursement_excel(reimbursements: List[StoreReimbursement]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Store Reimbursements"

    headers = [
        "Reimbursement No",
        "Store ID",
        "Store Name",
        "Region",
        "Channel",
        "Entity",
        "POC Name",
        "POC Phone",
        "Purchase Date",
        "Total Amount",
        "Paid Amount",
        "Outstanding Amount",
        "Status",
        "Reason",
        "Instructed By",
        "Created By",
        "Approved By",
        "Created At (UTC)",
        "Approved At (UTC)",
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 25

    for r in reimbursements:
        ws.append([
            r.reimbursement_number,
            r.store_id,
            r.store_name,
            r.store_region or "",
            r.store_channel or "",
            r.store_entity or "",
            r.poc_name or "",
            r.poc_phone or "",
            r.purchase_date.strftime("%Y-%m-%d") if r.purchase_date else "",
            float(r.total_amount),
            float(r.total_paid),
            float(r.outstanding_amount),
            r.status,
            r.reason,
            r.instructed_by_name or "",
            r.created_by.full_name if r.created_by else "",
            r.approved_by.full_name if r.approved_by else "",
            r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "",
            r.approved_at.strftime("%Y-%m-%d %H:%M:%S") if r.approved_at else "",
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
